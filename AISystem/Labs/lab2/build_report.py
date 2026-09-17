"""Создание русского отчёта по оформлению первой лабораторной работы."""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import pandas as pd
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT.parent / 'lab1' / 'Лабораторная_работа_1.docx'
OUTPUT = ROOT / 'Лабораторная_работа_6_Логистическая_регрессия.docx'
OUT = ROOT / 'output'
S = json.loads((OUT / 'summary.json').read_text(encoding='utf-8'))
GRID = pd.read_csv(OUT / 'hyperparameter_results.csv')
STAT = pd.read_csv(OUT / 'statistics_raw.csv', index_col=0)
LABELS = ['Беременности', 'Глюкоза', 'Давление', 'Кожная складка', 'Инсулин',
          'ИМТ', 'Наследственность', 'Возраст', 'Диабет']


def font(run, size=14, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold, run.font.italic = bold, italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')


def paragraph(doc, text, size=14, indent=True, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25) if indent and not center else Cm(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), size)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(text), bold=True, italic=level == 1)


def table(doc, headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment, t.autofit = 1, False
    if widths is None:
        widths = [17.05 / len(headers)] * len(headers)
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for k, row in enumerate([headers] + list(rows)):
        cells = t.rows[0].cells if k == 0 else t.add_row().cells
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if k == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        for i, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1
            font(p.add_run(str(value)), size, bold=k == 0)
            pr = cell._tc.get_or_add_tcPr()
            if k == 0 or k % 2 == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'D9E2F3' if k == 0 else 'F7F9FC')
                pr.append(shade)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{edge}')
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + name), value)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(doc, name, caption, width=6.45):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(OUT / 'figures' / name), width=Inches(width))
    p = paragraph(doc, caption, size=12, center=True)
    for r in p.runs:
        r.italic = True


def page(doc):
    doc.add_page_break()


def number(x, digits=4):
    return f'{x:.{digits}f}'.replace('.', ',')


def build():
    doc = Document(REFERENCE)
    cover_tables = [deepcopy(t._tbl) for t in doc.tables[:2]]
    # Сохраняем титул эталона; заменяем оглавление и основной текст.
    start = doc.paragraphs[24]._p
    body = doc._element.body
    removing = False
    for element in list(body):
        removing = removing or element is start
        if removing and element.tag != qn('w:sectPr'):
            body.remove(element)
    for i, text, size in [(8, 'ЛАБОРАТОРНАЯ РАБОТА 6', 18),
                          (9, 'ОТЧЕТ', 18), (10, 'по теме «Логистическая регрессия»', 14)]:
        p = doc.paragraphs[i]
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(text), size)
    paragraph(doc, 'Содержание', indent=False)
    p = doc.add_paragraph()
    run = p.add_run()
    begin, instruction, separate, end = [OxmlElement(t) for t in ['w:fldChar', 'w:instrText', 'w:fldChar', 'w:fldChar']]
    begin.set(qn('w:fldCharType'), 'begin')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate.set(qn('w:fldCharType'), 'separate')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, end])
    page(doc)

    heading(doc, 'ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ')
    paragraph(doc, 'Выбран набор данных Pima Indians Diabetes. Требуется обработать данные, получить числовую и графическую статистику, реализовать логистическую регрессию без готовых моделей и исследовать влияние гиперпараметров.')
    paragraph(doc, 'Сравниваются коэффициент обучения, число итераций и метод оптимизации. Для каждой комбинации на тестовой выборке рассчитываются доля верных ответов, точность, полнота и F1-мера. Допустимые сторонние зависимости вычислительной программы — только NumPy и Pandas.')
    heading(doc, 'ВВЕДЕНИЕ')
    paragraph(doc, 'Цель работы — самостоятельно реализовать вероятностный классификатор и исследовать его поведение на задаче распознавания диабета. Выполнены 40 экспериментов. Параметры выбираются по проверочной выборке, после чего модель повторно обучается на объединённых обучающих и проверочных данных.')
    paragraph(doc, 'В выбранной конфигурации используется градиентный спуск с шагом 0,01 и 100 итерациями. После повторного обучения доля верных ответов на тесте составляет 0,7532, F1-мера — 0,6275. Эти значения относятся к фиксированному разбиению и не доказывают универсального превосходства конфигурации.')
    heading(doc, 'ОПИСАНИЕ МЕТОДА')
    paragraph(doc, 'Логистическая регрессия моделирует вероятность положительного класса. Для каждого объекта вычисляется линейная комбинация признаков, затем применяется сигмоида. В отличие от линейной регрессии, результат всегда находится в интервале от 0 до 1.')
    paragraph(doc, 'zᵢ = b + xᵢᵀw;     pᵢ = σ(zᵢ) = 1 / (1 + exp(−zᵢ))', center=True)
    paragraph(doc, 'log(pᵢ / (1 − pᵢ)) = zᵢ;     ŷᵢ = 1, если pᵢ ≥ 0,5', center=True)

    page(doc)
    heading(doc, 'Функция потерь и оптимизация', 2)
    paragraph(doc, 'В соответствии с учебным материалом о логистической регрессии целевая переменная имеет распределение Бернулли. Максимизация правдоподобия эквивалентна минимизации средней бинарной перекрёстной энтропии:')
    paragraph(doc, 'L(θ) = −(1/n) Σᵢ [yᵢ ln pᵢ + (1 − yᵢ) ln(1 − pᵢ)]', center=True)
    paragraph(doc, 'В коде используется эквивалентное устойчивое выражение mean(logaddexp(0, z) − y·z). К матрице признаков добавляется столбец единиц; поэтому свободный член входит в общий вектор θ. Ниже X уже включает этот столбец.')
    paragraph(doc, 'g = Xᵀ(p − y) / n;     θ ← θ − αg', center=True)
    paragraph(doc, 'H = Xᵀ diag(pᵢ(1 − pᵢ)) X / n', center=True)
    paragraph(doc, '(H + 10⁻⁸I)δ = g;     θ ← θ − αδ', center=True)
    paragraph(doc, 'В методе Ньютона система решается функцией NumPy, но готовый классификатор не используется. Добавка 10⁻⁸I стабилизирует решение и не является штрафом в функции потерь. Если потери растут, шаг Ньютона уменьшается вдвое. Обе реализации выполняют заданное число итераций без досрочной остановки.')
    heading(doc, 'ПСЕВДОКОД МЕТОДА')
    for line in ['1. Загрузить таблицу и рассчитать описательную статистику.',
                 '2. Стратифицированно выделить обучение, проверку и тест.',
                 '3. По обучению определить медианы, средние и масштабы.',
                 '4. Для каждого метода, шага и числа итераций:',
                 '   4.1. Обнулить веса; вычислять вероятности и градиент.',
                 '   4.2. Обновлять веса выбранным методом оптимизации.',
                 '   4.3. Рассчитать четыре метрики на проверке и тесте.',
                 '5. Выбрать параметры по проверочной F1-мере.',
                 '6. Переобучить на объединении обучения и проверки.',
                 '7. Сохранить тестовые метрики, модель и графики.']:
        p = paragraph(doc, line, size=10.5, indent=False)
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            r.font.name = 'Courier New'

    page(doc)
    heading(doc, 'ПОДГОТОВКА ДАННЫХ')
    heading(doc, 'Описание набора данных', 2)
    paragraph(doc, 'Набор содержит 768 наблюдений, восемь числовых признаков и бинарную цель. В описании источника указано, что обследованы женщины народа пима в возрасте не менее 21 года. Отрицательный класс содержит 500 объектов, положительный — 268. Полных дубликатов и явных пустых ячеек нет.')
    table(doc, ['Поле', 'Смысл'], zip(STAT.index, ['Количество беременностей', 'Концентрация глюкозы', 'Диастолическое давление', 'Толщина кожной складки', 'Концентрация инсулина', 'Индекс массы тела', 'Функция наследственной предрасположенности', 'Возраст в годах', 'Диабет: 0 — нет, 1 — есть']), [6.4, 10.65], 10)
    paragraph(doc, 'В таблице приведена исходная статистика, включая нулевые значения. Обозначения: N — число непустых наблюдений, s — выборочное стандартное отклонение, Q1 и Q3 — квартили, Q2 — медиана. Во всех строках N = 768.')
    rows = [[label] + [number(row[c], 2) for c in ['mean', 'std', 'min', '25%', '50%', '75%', 'max']] for label, (_, row) in zip(LABELS, STAT.iterrows())]
    table(doc, ['Признак', 'Среднее', 's', 'Мин.', 'Q1', 'Q2', 'Q3', 'Макс.'], rows, [4.0] + [1.86] * 7, 9)
    paragraph(doc, 'Статистика после замены некорректных нулей на пропуски сохранена отдельно. Общая статистика служит только для описания и не используется для подбора параметров заполнения или стандартизации.', size=12)

    page(doc)
    heading(doc, 'Графическая описательная статистика', 2)
    figure(doc, 'statistics.png', 'Рисунок 1 — Количество, среднее, разброс, границы и квартили', width=6.1)
    paragraph(doc, 'Каждая панель соответствует одной статистике. Порядок столбцов совпадает с таблицей признаков; «Наслед.» означает функцию наследственной предрасположенности. Единицы разных признаков различаются, поэтому высота столбца не показывает важность признака.', size=12)
    paragraph(doc, 'Для инсулина среднее 79,80 заметно выше медианы 30,50, а максимум равен 846. Исходное распределение асимметрично и содержит много нулевых значений. Поэтому среднее по исходной таблице нельзя безоговорочно интерпретировать как типичный физиологический уровень.', size=12)

    page(doc)
    heading(doc, 'Предварительная обработка и разбиение', 2)
    paragraph(doc, 'Нули в глюкозе, давлении, кожной складке, инсулине и ИМТ рассматриваются как пропуски. Ноль беременностей имеет содержательный смысл и сохраняется. Количество заменённых значений приведено ниже.')
    table(doc, ['Признак', 'Число пропусков'], [('Глюкоза', 5), ('Давление', 35), ('Кожная складка', 227), ('Инсулин', 374), ('ИМТ', 11)], [10, 7.05])
    paragraph(doc, 'Сначала выделяются 614 объектов для разработки и 154 для тестирования — примерно 80/20. Из первой части 154 объекта выделяются для проверки. Итог: 460 обучающих, 154 проверочных и 154 тестовых объекта. Разбиение стратифицированное, генератор случайных чисел имеет начальное значение 42.')
    paragraph(doc, 'Медианы вычисляются только на обучающих данных. После заполнения пропусков из каждого признака вычитается обучающее среднее и результат делится на обучающее стандартное отклонение с делителем n. Те же параметры применяются к проверке и тесту. При повторном обучении параметры заново оцениваются на 614 объектах, без теста.')
    paragraph(doc, 'x′ⱼ = (xⱼ − μⱼ) / sⱼ', center=True)
    heading(doc, 'План исследования', 2)
    table(doc, ['Гиперпараметр', 'Проверенные значения'], [('Метод', 'Градиентный спуск; метод Ньютона'), ('Коэффициент обучения α', '0,001; 0,01; 0,1; 1'), ('Число итераций', '5; 20; 100; 500; 2000')], [7, 10.05])
    paragraph(doc, 'Всего проверены 2 × 4 × 5 = 40 комбинаций. По условию задания для всех комбинаций показаны тестовые метрики. Выбор конфигурации использует только максимальную проверочную F1-меру; при равенстве — меньшие проверочные потери, затем меньшее число итераций и фиксированный порядок. Тестовые результаты не участвуют в выборе.')

    page(doc)
    heading(doc, 'РЕЗУЛЬТАТЫ ВЫПОЛНЕНИЯ')
    heading(doc, 'Метрики и градиентный спуск', 2)
    paragraph(doc, 'Положительным считается класс диабета. TP и TN — верные положительные и отрицательные ответы, FP и FN — ложные положительные и отрицательные ответы. При нулевом знаменателе метрика принимается равной нулю.', size=12)
    paragraph(doc, 'Доля верных = (TP + TN) / N;  точность = TP / (TP + FP)', size=12, center=True)
    paragraph(doc, 'Полнота = TP / (TP + FN);  F1 = 2TP / (2TP + FP + FN)', size=12, center=True)
    headers = ['α', 'Итер.', 'Доля\nверных', 'Точность', 'Полнота', 'F1 тест', 'F1 пров.']
    widths = [1.45, 1.55, 2.75, 2.85, 2.85, 2.75, 2.85]
    for method in ['gd', 'newton']:
        if method == 'newton':
            page(doc)
            heading(doc, 'Метод Ньютона', 2)
            paragraph(doc, 'Во всех строках используются те же обучающие и тестовые объекты и тот же порог 0,5. Коэффициент α здесь задаёт длину шага вдоль направления Ньютона, поэтому его численное значение не эквивалентно шагу градиентного спуска.', size=12)
        rows = []
        for r in GRID[GRID.method == method].itertuples():
            rows.append([f'{r.rate:g}'.replace('.', ','), r.iterations] + [number(v) for v in [r.test_accuracy, r.test_precision, r.test_recall, r.test_f1, r.val_f1]])
        table(doc, headers, rows, widths, 10)
        if method == 'gd':
            paragraph(doc, 'Наибольшая проверочная F1-мера градиентного спуска равна 0,5794 при α = 0,01 и 100 итерациях. При α = 1 уже после 20 итераций тестовая F1-мера равна 0,6522, однако соответствующая проверочная F1-мера ниже — 0,5333. Поэтому по заранее заданному правилу выбирается первая конфигурация.', size=12)
        else:
            paragraph(doc, 'При α = 1 метод Ньютона быстро достигает практически постоянных весов: результаты для 20–2000 итераций совпадают с точностью таблицы. Малый шаг требует большего числа итераций. Полный CSV дополнительно содержит потери, TP, TN, FP и FN для каждой комбинации.', size=12)

    page(doc)
    heading(doc, 'Влияние гиперпараметров на качество', 2)
    figure(doc, 'gd_metrics.png', 'Рисунок 2 — Градиентный спуск и четыре тестовые метрики', width=6.0)
    paragraph(doc, 'Горизонтальная ось показывает число итераций в логарифмическом масштабе, цвет — коэффициент обучения. Малые шаги изменяют веса медленно. Рост числа итераций не гарантирует монотонного улучшения полноты и F1-меры: предсказанные классы меняются дискретно при пересечении порога 0,5.')
    figure(doc, 'newton_metrics.png', 'Рисунок 3 — Метод Ньютона и четыре тестовые метрики', width=6.0)

    page(doc)
    heading(doc, 'Сходимость и итоговая модель', 2)
    figure(doc, 'loss.png', 'Рисунок 4 — Обучающие потери в зависимости от числа итераций')
    paragraph(doc, 'При α = 1 после 20 итераций градиентный спуск даёт потери 0,449029, а метод Ньютона — 0,446514. После 2000 итераций оба метода достигают 0,446514. Ньютон требует меньше итераций, но каждая итерация включает построение матрицы Гессе и решение системы.')
    paragraph(doc, 'У выбранной по проверочной F1 конфигурации обучающие потери равны 0,581274: она ещё не сошлась. Ограничение числа шагов действует подобно ранней остановке. Разница проверочной F1 с ближайшей альтернативой составляет лишь 0,002516; считать её статистически значимой на одной выборке нельзя.')
    m = S['final']
    table(doc, ['Модель', 'Доля верных', 'Точность', 'Полнота', 'F1'],
          [['Всегда класс 0'] + [number(S['baseline'][k]) for k in ['accuracy', 'precision', 'recall', 'f1']],
           ['После переобучения'] + [number(m[k]) for k in ['accuracy', 'precision', 'recall', 'f1']]],
          [5.05, 3, 3, 3, 3])
    table(doc, ['Истинный класс', 'Предсказан 0', 'Предсказан 1'], [['0', m['tn'], m['fp']], ['1', m['fn'], m['tp']]], [7.05, 5, 5])
    paragraph(doc, f'После обучения на 614 объектах тестовые потери равны {number(m["log_loss"], 6)}. Модель верно распознала 32 положительных случая и пропустила 22. Доля верных ответов выше базового уровня 0,6494, но полнота 0,5926 остаётся ограниченной.', size=12)

    page(doc)
    heading(doc, 'ПРИМЕРЫ ИСПОЛЬЗОВАНИЯ МЕТОДА')
    paragraph(doc, 'Модель может использоваться как учебный пример бинарного классификатора: по восьми измерениям вычисляются вероятность класса 1 и ответ при пороге 0,5. Аналогичная схема применима к классификации сообщений или оценке вероятности события, если доступны размеченные данные.')
    paragraph(doc, 'Вектор весов и параметры заполнения и стандартизации сохранены в model.npz. При применении к новым объектам необходимо сохранить порядок признаков и использовать эти параметры, а не пересчитывать их на новых данных. Коэффициент exp(wⱼ) характеризует изменение шансов при увеличении стандартизованного признака на единицу, а не изменение вероятности в столько же раз.')
    heading(doc, 'ЗАКЛЮЧЕНИЕ')
    paragraph(doc, 'Реализованы сигмоида, устойчивая функция потерь, пакетный градиентный спуск и метод Ньютона. Получены числовые и графические характеристики данных, обработаны неявные пропуски и выполнено разделение без утечки параметров предобработки. Все 40 комбинаций оценены четырьмя метриками.')
    paragraph(doc, 'По принятому правилу выбран градиентный спуск с α = 0,01 и 100 итерациями. Итоговые тестовые показатели: доля верных ответов 0,7532; точность 0,6667; полнота 0,5926; F1-мера 0,6275. Метод Ньютона быстрее сходится по числу итераций, однако лучшая сходимость обучающей функции не означает наилучшую проверочную F1-меру.')
    paragraph(doc, 'Выводы ограничены одной небольшой выборкой и одним разбиением. Тестовые показатели всех конфигураций опубликованы по условию задания, поэтому процедура не является полностью слепым тестированием. Для более надёжного выбора следует использовать стратифицированную перекрёстную проверку и затем новые независимые данные. Результаты не обосновывают перенос на другие группы населения.')
    heading(doc, 'Источники и воспроизводимость', 2)
    paragraph(doc, 'Учебный материал: «Модели логистической регрессии», глава 14, файл «логистическая_регрессия.docx». Набор данных: https://www.kaggle.com/datasets/uciml/pima-indians-diabetes-database', size=11, indent=False)
    paragraph(doc, 'Интерфейс загрузки Kaggle вернул код 403. Фактически использовано открытое зеркало: https://raw.githubusercontent.com/jbrownlee/Datasets/master/pima-indians-diabetes.data.csv. Побайтовое сравнение с недоступной выгрузкой Kaggle не выполнялось.', size=11, indent=False)
    paragraph(doc, 'Запуск: python logistic_regression.py. Зависимости: NumPy и Pandas. Файл graphics.py строит SVG средствами стандартной библиотеки. Таблицы, метрики, разбиение и контрольная сумма данных находятся в output; программа build_report.py отвечает только за оформление Word.', size=11, indent=False)

    authored = ROOT / '_qa' / 'authored.docx'
    doc.save(authored)
    # Сохраняем непричастные части пакета исходного шаблона побайтно.
    with ZipFile(REFERENCE) as src, ZipFile(authored) as new, ZipFile(OUTPUT, 'w', ZIP_DEFLATED) as dst:
        editable = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
        for name in src.namelist():
            dst.writestr(name, new.read(name) if name in editable else src.read(name))
        for name in new.namelist():
            if name not in src.namelist():
                dst.writestr(name, new.read(name))
    result = Document(OUTPUT)
    assert all(a.xml == b._tbl.xml for a, b in zip(cover_tables, result.tables[:2]))
    print(OUTPUT)


if __name__ == '__main__':
    build()
