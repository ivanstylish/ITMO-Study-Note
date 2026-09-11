from pathlib import Path
from copy import deepcopy
import os

ROOT = Path(__file__).resolve().parent
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".matplotlib"))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


TEMPLATE = ROOT.parent / "lab1.docx"
OUTPUT = ROOT / "Лабораторная_работа_1_Линейная_регрессия.docx"
FIGURES = ROOT / "output" / "figures"
METRICS = pd.read_csv(ROOT / "output" / "model_metrics.csv")
COEFFICIENTS = pd.read_csv(ROOT / "output" / "coefficients.csv")
STATS = pd.read_csv(ROOT / "output" / "summary_statistics.csv", index_col=0)


def set_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True, after=5):
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.first_line_indent = Cm(1.25) if first_line else None
    for run in paragraph.runs:
        set_font(run)


def add_text(doc, text, bold_start=None):
    paragraph = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        first = paragraph.add_run(bold_start)
        set_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_start):])
        set_font(rest)
    else:
        set_font(paragraph.add_run(text))
    style_paragraph(paragraph)
    return paragraph


def add_heading(doc, text, level=1):
    style_name = f"Heading {level}"
    if style_name not in [style.name for style in doc.styles]:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level - 1))
        style._element.get_or_add_pPr().append(outline)
    paragraph = doc.add_paragraph(style=style_name)
    run = paragraph.add_run(text)
    set_font(run, size=14, bold=True, italic=(level == 1))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), "70")
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_table_borders(table, color="D9D9D9"):
    table_pr = table._tbl.tblPr
    borders = table_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, size=font_size)
        shade_cell(table.rows[0].cells[i])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for row_number, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell(cells[i], value, align=align, size=font_size)
            if row_number % 2 == 1:
                shade_cell(cells[i], "F7F9FC")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    set_table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, path, caption, width=6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))

    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.first_line_indent = None
    caption_p.paragraph_format.space_after = Pt(7)
    caption_run = caption_p.add_run(caption)
    set_font(caption_run, size=12, italic=True)


def add_toc(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("Содержание"), size=14)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Содержание будет обновлено при открытии документа"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_font(run, size=12)


def create_equation_image():
    path = FIGURES / "03_equations.png"
    fig = plt.figure(figsize=(10, 2.2), facecolor="white")
    lines = [
        r"$\hat{y}=\theta_0+\sum_{j=1}^{p}\theta_j x_j$",
        r"$J(\theta)=\frac{1}{n}\sum_{i=1}^{n}(\hat{y}_i-y_i)^2$",
        r"$\theta\leftarrow\theta-\eta\frac{2}{n}X^{T}(X\theta-y)$",
    ]
    for y, line in zip([0.78, 0.48, 0.18], lines):
        fig.text(0.5, y, line, ha="center", va="center", fontsize=19)
    plt.axis("off")
    fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    return path


def remove_template_body(doc):
    body = doc._element.body
    removing = False
    for child in list(body):
        if child.tag == qn("w:sdt"):
            removing = True
        if removing and child.tag != qn("w:sectPr"):
            body.remove(child)


def update_title_page(doc):
    replacements = {
        8: ("ЛАБОРАТОРНАЯ РАБОТА 1", 18, False),
        9: ("ОТЧЕТ", 18, False),
        10: ("по теме «Линейная регрессия»", 14, False),
    }
    for index, (text, size, italic) in replacements.items():
        paragraph = doc.paragraphs[index]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        set_font(run, size=size, italic=italic)


def set_update_fields(doc):
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def build():
    doc = Document(TEMPLATE)
    update_title_page(doc)
    remove_template_body(doc)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(14)

    add_toc(doc)
    doc.add_page_break()

    add_heading(doc, "ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ")
    add_text(doc, "Для варианта с нечетным номером используется набор данных об учебной успеваемости студентов Student_Performance.csv.")
    tasks = [
        "получить числовую и графическую описательную статистику, включая среднее, стандартное отклонение, минимум, максимум и квартильные значения;",
        "обработать пропуски, категориальные признаки и выполнить нормализацию;",
        "разделить данные на обучающую и тестовую выборки;",
        "реализовать линейную регрессию методом наименьших квадратов без готовых библиотечных моделей;",
        "построить три модели с различными наборами признаков и оценить их коэффициентом детерминации R²;",
        "сравнить модели и определить наиболее полезные признаки;",
        "в рамках бонусного задания добавить синтетический признак.",
    ]
    for task in tasks:
        p = doc.add_paragraph()
        set_font(p.add_run("• "), bold=True)
        set_font(p.add_run(task))
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(3)

    add_heading(doc, "ВВЕДЕНИЕ")
    add_text(doc, "Цель работы состоит в практической реализации множественной линейной регрессии и проверке влияния различных групп признаков на прогноз интегрального показателя успеваемости студентов.")
    add_text(doc, "Для достижения цели выполнены разведочный анализ данных, очистка и кодирование признаков, стандартизация, разбиение выборки, обучение трех моделей собственным алгоритмом градиентного спуска и сравнение результатов на независимой тестовой части данных.")

    add_heading(doc, "ОПИСАНИЕ МЕТОДА")
    add_text(doc, "Линейная регрессия описывает зависимость целевой переменной от набора признаков линейной комбинацией. Параметры модели подбираются так, чтобы средний квадрат разности между фактическими и предсказанными значениями был минимальным.")
    equation = create_equation_image()
    add_figure(doc, equation, "Рисунок 1 - Модель, функция потерь и шаг градиентного спуска", width=6.0)
    add_text(doc, "В работе коэффициенты не вычисляются готовой функцией регрессии или решателем системы линейных уравнений. Они последовательно обновляются пакетным градиентным спуском. Перед обучением признаки стандартизируются, поэтому коэффициенты внутри одной модели можно сравнивать по модулю.")
    add_text(doc, "Качество оценивается коэффициентом детерминации R². Значение, близкое к 1, означает, что модель объясняет большую долю изменчивости целевой переменной. Дополнительно рассчитаны RMSE и MAE в единицах индекса успеваемости.")

    add_heading(doc, "ПСЕВДОКОД МЕТОДА")
    pseudo = [
        "Вход: таблица D, целевой признак y, три набора признаков F1, F2, F3",
        "1. Удалить полные дубликаты строк.",
        "2. Заполнить пропуски: медианой для чисел и модой для категорий.",
        "3. Закодировать категориальные признаки и создать синтетический признак.",
        "4. Случайно разделить D на 80% обучающих и 20% тестовых наблюдений.",
        "5. Для каждого набора F:",
        "   5.1. Стандартизировать признаки по параметрам обучающей выборки.",
        "   5.2. Инициализировать вектор коэффициентов нулями.",
        "   5.3. Повторять: прогноз -> ошибка -> градиент -> обновление коэффициентов.",
        "   5.4. Остановиться при малом изменении коэффициентов.",
        "   5.5. Получить прогнозы и вычислить R², RMSE и MAE.",
        "6. Сравнить три модели и выбрать модель с максимальным R².",
        "Выход: коэффициенты, метрики и графики.",
    ]
    for line in pseudo:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(line), name="Courier New", size=10.5)

    add_heading(doc, "ПОДГОТОВКА ДАННЫХ")
    add_heading(doc, "Описание набора данных", level=2)
    add_text(doc, "Исходный файл содержит 10 000 наблюдений и 6 признаков. Целевой столбец Performance Index принимает значения от 10 до 100. После удаления 127 полных дубликатов осталось 9 873 уникальных наблюдения. Пропущенные значения отсутствуют.")
    rows = []
    labels = {
        "Hours Studied": "Часы учебы",
        "Previous Scores": "Предыдущие баллы",
        "Sleep Hours": "Часы сна",
        "Sample Question Papers Practiced": "Решенные варианты",
        "Performance Index": "Индекс успеваемости",
    }
    for feature, label in labels.items():
        row = STATS.loc[feature]
        rows.append([
            label,
            f"{row['mean']:.2f}", f"{row['std']:.2f}", f"{row['min']:.0f}",
            f"{row['25%']:.0f}", f"{row['50%']:.0f}", f"{row['75%']:.0f}", f"{row['max']:.0f}",
        ])
    add_table(
        doc,
        ["Признак", "Среднее", "Ст. откл.", "Мин.", "Q1", "Медиана", "Q3", "Макс."],
        rows,
        widths=[1.55, 0.72, 0.78, 0.55, 0.5, 0.72, 0.5, 0.55],
        font_size=9,
    )
    add_figure(doc, FIGURES / "01_eda.png", "Рисунок 2 - Описательная статистика и связи между признаками", width=6.25)
    add_text(doc, "Наиболее сильная линейная связь с целевой переменной наблюдается у Previous Scores: коэффициент корреляции равен 0,915. Для Hours Studied корреляция составляет 0,375. Связь часов сна и числа решенных вариантов с целевым индексом значительно слабее. Категории внеучебной активности почти сбалансированы: 4 986 значений No и 4 887 значений Yes.")

    add_heading(doc, "Предварительная обработка", level=2)
    add_text(doc, "Категориальный признак Extracurricular Activities преобразован в бинарный индикатор. Для потенциальных пропусков предусмотрено заполнение медианой или модой. Все признаки стандартизируются с использованием среднего и стандартного отклонения, рассчитанных только на обучающей выборке. Выборка с фиксированным seed 42 разделена на 7 898 обучающих и 1 975 тестовых наблюдений.")
    add_text(doc, "Синтетический признак Study Efficiency определен как произведение Hours Studied и Previous Scores, деленное на 100. Он отражает совместное влияние интенсивности учебы и исходного уровня подготовки.")

    add_heading(doc, "РЕЗУЛЬТАТЫ ВЫПОЛНЕНИЯ")
    model_rows = [
        ["Модель 1", "Учебный режим", "Hours Studied, Sleep Hours, Sample Question Papers Practiced"],
        ["Модель 2", "Исходная подготовка", "Previous Scores, Extracurricular Activities"],
        ["Модель 3", "Полный набор", "Все исходные признаки и Study Efficiency"],
    ]
    add_table(doc, ["Модель", "Смысл набора", "Признаки"], model_rows, widths=[0.9, 1.45, 4.1], font_size=9.5)

    metric_rows = []
    for _, row in METRICS.iterrows():
        metric_rows.append([row["model"], f"{row['R2']:.4f}", f"{row['RMSE']:.4f}", f"{row['MAE']:.4f}"])
    add_table(doc, ["Модель", "R²", "RMSE", "MAE"], metric_rows, widths=[1.6, 1.2, 1.2, 1.2], font_size=10)
    add_figure(doc, FIGURES / "02_models.png", "Рисунок 3 - Сравнение качества моделей и прогноз лучшей модели", width=6.25)

    add_text(doc, "Модель 1 объясняет только 14,68% вариации индекса успеваемости. Одних показателей режима учебы и сна недостаточно для точного прогноза. Модель 2 достигает R² = 0,8377 благодаря признаку Previous Scores, который отражает исходный уровень знаний. Модель 3 использует полный набор и показывает лучший результат: R² = 0,9886, RMSE = 2,0648 и MAE = 1,6398.")
    add_text(doc, "В третьей модели наибольшие стандартизованные коэффициенты имеют Previous Scores (17,6080) и Hours Studied (7,3571). Вклад Sleep Hours равен 0,8213, числа решенных вариантов - 0,5489, внеучебной активности - 0,3104, а синтетического Study Efficiency - 0,0179. Следовательно, главные факторы качества прогноза - предыдущие результаты и продолжительность учебы. Синтетический признак реализует бонусное условие, но после учета исходных признаков дает небольшой дополнительный линейный вклад.")

    add_heading(doc, "ПРИМЕРЫ ИСПОЛЬЗОВАНИЯ МЕТОДА")
    add_text(doc, "Линейная регрессия может использоваться для предварительного прогноза итоговой успеваемости по текущим учебным показателям. Такой прогноз помогает выявить студентов, которым может потребоваться дополнительная поддержка, если решение принимается вместе с педагогической оценкой, а не только по числу модели.")
    add_text(doc, "Метод также подходит для анализа факторов результата: знак и величина коэффициентов показывают направление и относительную силу связи признаков с целевой переменной. За пределами образования линейная регрессия применяется для прогноза спроса, затрат, времени выполнения и других непрерывных величин, если зависимость приблизительно линейна и данные сопоставимы.")

    add_heading(doc, "ЗАКЛЮЧЕНИЕ")
    add_text(doc, "Эксперимент показал, что качество линейной регрессии определяется информативностью выбранных признаков. Модель учебного режима без данных о прошлой успеваемости имеет низкий R² = 0,1468. Добавление Previous Scores повышает показатель до 0,8377. Полная модель достигает R² = 0,9886 и средней абсолютной ошибки 1,6398 пункта.")
    add_text(doc, "Полученные результаты подтверждают, что Previous Scores и Hours Studied являются наиболее полезными признаками в данном наборе. Внеучебная активность, сон, число решенных вариантов и синтетический показатель дают меньший вклад. Реализация градиентного спуска без готовой функции регрессии корректно находит коэффициенты, минимизирующие квадратичную ошибку на обучающей выборке.")

    set_update_fields(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
