"""Оформление отчёта из фактически сохранённых результатов эксперимента."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree

from regularization import ROOT, MODELS, VARIANTS

REFERENCE = ROOT / "report.docx"
FINAL = ROOT / "Лабораторная_работа_3_Регуляризация.docx"
QA = ROOT / "_qa"
EDITABLE = {"word/document.xml", "word/_rels/document.xml.rels", "word/styles.xml",
            "word/settings.xml", "[Content_Types].xml"}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
SHORT = {"Hours Studied": "Часы учёбы", "Previous Scores": "Предыдущие баллы",
         "Sleep Hours": "Часы сна", "Sample Question Papers Practiced": "Решённые варианты",
         "Extracurricular Activities": "Внеучебная активность", "Study Interaction": "Взаимодействие"}


def add_styles(doc):
    for name in ["Lab Body", "Lab Caption", "Lab Code", "Title"]:
        if name not in doc.styles:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles["Normal"]
    for name in ["Lab Body", "Lab Caption", "Lab Code", "Title"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman" if name != "Lab Code" else "Consolas"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12 if name == "Lab Body" else 10 if name != "Title" else 14)
        style.font.bold = name == "Title"
        style.font.italic = False
        pf = style.paragraph_format
        pf.line_spacing = 1.25 if name == "Lab Body" else 1.0
        pf.space_after = Pt(5 if name == "Lab Body" else 3)
        pf.space_before = Pt(0)
        pf.first_line_indent = Inches(0.5) if name == "Lab Body" else Cm(0)
        pf.widow_control = True


def paragraph(doc, text, style="Lab Body"):
    p = doc.add_paragraph(text, style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style == "Lab Body" else WD_ALIGN_PARAGRAPH.LEFT
    return p


def heading(doc, text, new_page=True):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.1
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.italic = False
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def caption(doc, text):
    p = paragraph(doc, text, "Lab Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    return p


def image(doc, path, text, width=17):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    caption(doc, text)


def table(doc, headers, rows, widths, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for row_index, values in enumerate([headers] + list(rows)):
        cells = t.rows[0].cells if row_index == 0 else t.add_row().cells
        trpr = t.rows[row_index]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
        for column, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                tag = OxmlElement("w:" + edge)
                tag.set(qn("w:w"), "80")
                tag.set(qn("w:type"), "dxa")
                margins.append(tag)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                tag = OxmlElement("w:" + edge)
                for attribute, val in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    tag.set(qn("w:" + attribute), val)
                borders.append(tag)
            tcpr.append(borders)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7E6E6")
                tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.font.bold = row_index == 0
            run.font.color.rgb = RGBColor(0, 0, 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.line_spacing = 0.3
    after.add_run().font.size = Pt(2)
    return t


def mr(text):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = text
    r.append(t)
    return r


def fraction(numerator, denominator):
    f = OxmlElement("m:f")
    for tag, content in [("m:num", numerator), ("m:den", denominator)]:
        part = OxmlElement(tag)
        part.append(mr(content))
        f.append(part)
    return f


def equation(doc, *parts):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    math = OxmlElement("m:oMath")
    for part in parts:
        math.append(mr(part) if isinstance(part, str) else part)
    p._p.append(math)


def replace_package(source, destination):
    """Сохраняем неизменяемые части шаблона побайтно."""
    with ZipFile(REFERENCE) as original, ZipFile(source) as authored:
        original_names = set(original.namelist())
        authored_names = set(authored.namelist())
        with ZipFile(destination, "w", ZIP_DEFLATED) as final:
            for name in original.namelist():
                final.writestr(name, authored.read(name) if name in EDITABLE else original.read(name))
            for name in authored.namelist():
                if name not in original_names:
                    final.writestr(name, authored.read(name))
    with ZipFile(REFERENCE) as original, ZipFile(destination) as final:
        for name in original.namelist():
            if name not in EDITABLE and original.read(name) != final.read(name):
                raise AssertionError("Изменилась защищённая часть: " + name)


def create_report():
    out = ROOT / "output"
    s = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    if (s["seed"] != 42 or s["learning_rate"] != 0.05 or s["iterations"] != 10000
            or s["lambdas"] != [0.001, 0.01, 0.1, 1.0, 10.0]
            or s["l1_ratios"] != [0.2, 0.5, 0.8] or s["clean_rows"] != 9873):
        raise ValueError("Аналитический текст отчёта описывает стандартный запуск. "
                         "После изменения данных/параметров обновите текст build_report.py.")
    metrics = pd.read_csv(out / "metrics.csv")
    coeff = pd.read_csv(out / "coefficients.csv")
    grid = pd.read_csv(out / "hyperparameter_results.csv")
    doc = Document(REFERENCE)
    body = doc._element.body
    # Сохраняем весь титульный лист и настройки секции, заменяем старое оглавление.
    for child in list(body)[26:-1]:
        body.remove(child)
    add_styles(doc)
    cover_slot = body[11]
    subject = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    props.append(fonts)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), "28")
    props.append(font_size)
    subject.append(props)
    subject_text = OxmlElement("w:t")
    subject_text.text = "Регуляризация моделей"
    subject.append(subject_text)
    cover_slot.append(subject)
    cover_props = cover_slot.find(qn("w:pPr"))
    if cover_props is None:
        cover_props = OxmlElement("w:pPr")
        cover_slot.insert(0, cover_props)
    title_style = OxmlElement("w:pStyle")
    title_style.set(qn("w:val"), "Title")
    cover_props.insert(0, title_style)
    align = OxmlElement("w:jc")
    align.set(qn("w:val"), "center")
    cover_props.append(align)

    p = doc.add_paragraph("Содержание", style="Lab Body")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.first_line_indent = Cm(0)
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.bold = True
    toc = doc.add_paragraph()
    for tag, attribute in [("w:fldChar", "begin"), ("w:instrText", None),
                           ("w:fldChar", "separate"), ("w:fldChar", "end")]:
        run = OxmlElement("w:r")
        element = OxmlElement(tag)
        if attribute:
            element.set(qn("w:fldCharType"), attribute)
        else:
            element.text = ' TOC \\o "1-1" \\h \\z \\u '
            element.set(qn("xml:space"), "preserve")
        run.append(element)
        toc._p.append(run)

    heading(doc, "ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ И ДАННЫЕ")
    paragraph(doc, "Цель работы — исследовать влияние L1, L2 и ElasticNet на три модели линейной регрессии из лабораторной работы 1. Сравниваются качество прогноза, разрыв между обучением и валидацией и значения коэффициентов. На полном наборе признаков получено тестовое R² около 0,9886. Слабая регуляризация почти не меняет качество.")
    paragraph(doc, "По заданию необходимо разделить данные на три выборки, рассчитывать нормировку только по обучающей части, реализовать штрафы без готовых моделей, не штрафовать свободный член, сравнить четыре варианта при одинаковых условиях, подобрать параметры на валидации, построить четыре кривые для каждой модели и оценить результаты на тесте.")
    paragraph(doc, f"Использован исходный Student_Performance.csv из лабораторной 1: {s['raw_rows']:,} строк, четыре числовых предиктора, один категориальный признак и целевой Performance Index. После удаления {s['duplicates']} полных дубликатов осталось {s['clean_rows']:,} наблюдений. Пропусков и строк без целевого значения в данном файле нет.")
    table(doc, ["Модель", "Состав признаков"], [
        ["Модель 1", "Hours Studied, Sleep Hours, Sample Question Papers Practiced"],
        ["Модель 2", "Previous Scores, Extracurricular Activities"],
        ["Модель 3", "Четыре числовых признака, Extracurricular Activities, Study Interaction"],
    ], [2.4, 14.6], size=11)
    paragraph(doc, "Признак Study Interaction равен Hours Studied × Previous Scores / 100. Это та же формула, которая используется в лабораторной 1. В её отчёте признак назван Study Efficiency, а в действующем коде — Study Interaction. Его значения и смысл сохраняются.")
    paragraph(doc, "Категории No и Yes кодируются числами 0 и 1. Для потенциальных пропусков предусмотрены медиана числового столбца и мода категории, найденные на обучающей выборке. После создания взаимодействия все предикторы, включая бинарный и синтетический, стандартизируются по обучающим среднему и стандартному отклонению с ddof = 0.")

    heading(doc, "МЕТОД И ОБУЧЕНИЕ")
    paragraph(doc, "Модель вычисляет линейную комбинацию стандартизованных признаков и свободного члена. Ошибка и качество для каждой выборки определяются следующим образом.")
    equation(doc, "ŷ = b + zᵀw")
    equation(doc, "MSE = ", fraction("1", "n"), " ∑ᵢ (ŷᵢ − yᵢ)²     R² = 1 − ", fraction("∑ᵢ (yᵢ − ŷᵢ)²", "∑ᵢ (yᵢ − ȳ)²"))
    paragraph(doc, "Во всех вариантах минимизируется одна общая функция. В суммах штрафа участвуют только коэффициенты предикторов w. Свободный член b исключён.")
    equation(doc, "J = MSE + λ [ρ ∑ⱼ |wⱼ| + ", fraction("1 − ρ", "2"), " ∑ⱼ wⱼ²]")
    paragraph(doc, "Без регуляризации λ = 0. Для L1 используется ρ = 1, для L2 — ρ = 0, для ElasticNet — 0 < ρ < 1. Здесь ρ обозначает долю L1 в штрафе. Множитель 1/2 в L2 определяет численный масштаб λ.")
    paragraph(doc, "Обучение выполнено полным пакетным градиентным спуском. Для негладкой L1-части после шага по MSE и L2 применяется мягкий порог. Это проксимальный вариант градиентного спуска, который корректно обрабатывает точку ноль и позволяет получить строго нулевые веса.")
    equation(doc, "vⱼ = wⱼ − η [∇ⱼMSE + λ(1 − ρ)wⱼ]")
    equation(doc, "wⱼ = sign(vⱼ) max(|vⱼ| − ηλρ, 0)")
    paragraph(doc, "Свободный член обновляется только по градиенту MSE. На кривых и в таблицах приведена чистая MSE без штрафа. Для ускорения заранее вычисляются XᵀX/n и Xᵀy/n. Они дают тот же полный градиент, что и прямой проход по всем объектам на каждом шаге.")
    for line in ["1. Очистить данные и создать общее разбиение 60/20/20.",
                 "2. Оценить заполнение пропусков и нормировку только на train.",
                 "3. Для каждого набора признаков обучить все комбинации λ и ρ.",
                 "4. Внутри каждого варианта выбрать минимум валидационной MSE.",
                 "5. Повторить выбранные четыре обучения с записью всех итераций.",
                 "6. Один раз оценить выбранные варианты на независимом тесте."]:
        paragraph(doc, line, "Lab Code")

    heading(doc, "ПРОТОКОЛ И ВЫБОР ПАРАМЕТРОВ")
    paragraph(doc, f"Фиксирован seed = {s['seed']}. Обучающая часть содержит {s['split_sizes']['train']} объектов, валидационная и тестовая — по {s['split_sizes']['test']}. Все три модели и четыре варианта используют эти же строки. Начальные веса равны нулю, шаг η = {s['learning_rate']}, число обновлений = {s['iterations']:,}. Досрочная остановка отключена. История содержит также начальную точку 0.")
    paragraph(doc, "Проверены пять значений λ от 0,001 до 10 с десятикратным увеличением каждого следующего значения. Для ElasticNet проверены доли L1, равные 0,2, 0,5 и 0,8. На одну модель приходится 1 обучение без штрафа, 5 с L1, 5 с L2 и 15 с ElasticNet. Всего 78 комбинаций. Тестовые метрики не использовались при выборе параметров.")
    rows = []
    for _, row in metrics.iterrows():
        rows.append([row.model[-1], row.variant, f"{row['lambda']:g}", f"{row.l1_ratio:g}", f"{row.validation_MSE:.6f}"])
    caption(doc, "Таблица 1 — Параметры выбранных вариантов и валидационная MSE")
    table(doc, ["№", "Вариант", "λ", "ρ", "MSE валидации"], rows, [1.4, 5.2, 2, 2, 6.4], size=11)
    paragraph(doc, "В модели 1 небольшой штраф улучшает валидационную ошибку, но выбранные варианты не превосходят исходный на тесте. Для модели 2 изменения малы. В модели 3 лучшая положительная λ лежит на нижней границе сетки. Исходная модель без штрафа имеет меньшую валидационную ошибку, чем все её регуляризованные варианты.")
    paragraph(doc, "После выбора параметров обучение повторено на той же обучающей части с теми же условиями. Объединение train и validation не выполнялось, поэтому итоговые кривые, коэффициенты и тестовые метрики описывают одни и те же обученные модели. Полная сетка сохранена в hyperparameter_results.csv.")

    heading(doc, "ВЛИЯНИЕ ПАРАМЕТРОВ РЕГУЛЯРИЗАЦИИ")
    paragraph(doc, "На рисунках показаны все пять значений λ и три доли L1 для ElasticNet. Логарифмическая ось λ позволяет видеть весь исследованный диапазон.")
    for number in range(1, 4):
        image(doc, out / "figures" / f"model_{number}_parameters.png",
              f"Рисунок {number} — Валидационная MSE модели {number} при изменении λ и ρ")
    paragraph(doc, "Сильный штраф вызывает недообучение: при λ = 10 в модели 3 валидационная MSE для L1 равна 42,8100, а для L2 — 222,3922 против 4,3044 без регуляризации. При этой же λ увеличение ρ с 0,2 до 0,8 уменьшает ошибку ElasticNet с 210,1585 до 122,6756, но она остаётся намного выше исходной.")

    curve_comments = [
        "Модель 1 не содержит Previous Scores. Её высокая ошибка объясняется недостатком информативных признаков. Регуляризация уменьшает нормы коэффициентов, однако не компенсирует отсутствие исходной подготовки. Валидационное улучшение L1 не переносится на тест.",
        "Модель 2 использует Previous Scores и значительно точнее первой. Кривые при выбранном слабом штрафе почти совпадают. L2 и ElasticNet дают очень небольшое улучшение на тесте. По одному разбиению нельзя считать такую разницу устойчивым преимуществом.",
        "Модель 3 использует полный набор и даёт наилучшее качество среди рассматриваемых наборов. Обучающая и валидационная ошибки стабилизируются без позднего роста валидационной MSE. При выбранной λ = 0,001 изменения качества очень малы.",
    ]
    for number, (name, _) in enumerate(MODELS.items(), 1):
        heading(doc, f"КРИВЫЕ ОБУЧЕНИЯ МОДЕЛИ {number}")
        paragraph(doc, "Для четырёх вариантов отображены MSE и R² на train и validation. Ось итераций использует симметричную логарифмическую шкалу, MSE — логарифмическую. Вставки увеличивают последние 10% итераций, чтобы различить близкие линии.")
        image(doc, out / "figures" / f"model_{number}_curves.png",
              f"Рисунок {number+3} — Четыре кривые качества для модели {number}")
        rows = [[r.variant, f"{r.train_MSE:.6f}", f"{r.validation_MSE:.6f}",
                 f"{r.train_R2:.6f}", f"{r.validation_R2:.6f}"] for _, r in metrics[metrics.model == name].iterrows()]
        table(doc, ["Вариант", "MSE train", "MSE val", "R² train", "R² val"], rows, [5, 3, 3, 3, 3], size=10)
        paragraph(doc, curve_comments[number-1])

    heading(doc, "ОЦЕНКА НА ТЕСТЕ И ПЕРЕОБУЧЕНИЕ")
    paragraph(doc, "На тестовой выборке из 1975 объектов оценены все 12 вариантов. MSE рассчитана без штрафа. RMSE и MAE выражены в единицах индекса успеваемости.")
    rows = [[r.model[-1], r.variant, f"{r.test_MSE:.6f}", f"{r.test_R2:.6f}",
             f"{r.test_RMSE:.4f}", f"{r.test_MAE:.4f}"] for _, r in metrics.iterrows()]
    caption(doc, "Таблица 2 — Метрики на независимой тестовой выборке")
    table(doc, ["№", "Вариант", "MSE", "R²", "RMSE", "MAE"], rows, [1.3, 4.4, 3.3, 3, 2.5, 2.5], size=10)
    paragraph(doc, "Без штрафа R² равен 0,146505 для модели 1, 0,837623 для модели 2 и 0,988622 для модели 3. Для модели 2 ElasticNet уменьшает MSE лишь с 60,823735 до 60,819923. Для полной модели все выбранные штрафы немного повышают MSE.")
    rows = []
    for name in MODELS:
        sub = metrics[metrics.model == name]
        rows.append([name] + [f"{sub[sub.variant == v].iloc[0].generalization_gap:.6f}" for v in VARIANTS])
    caption(doc, "Таблица 3 — Разрыв MSE validation минус MSE train")
    table(doc, ["Модель", "Без штрафа", "L1", "L2", "ElasticNet"], rows, [3, 3.5, 3.5, 3.5, 3.5], size=11)
    paragraph(doc, "У модели 1 разрыв уменьшается с ростом обучающей ошибки, поэтому это не доказывает улучшения тестового прогноза. Для модели 2 валидационная ошибка немного меньше обучающей. У модели 3 разрыв около 0,20 при MSE около 4. Выраженного переобучения по кривым не видно.")
    paragraph(doc, "Одно фиксированное разбиение не позволяет считать малые различия статистически значимым улучшением. Для всех вариантов сохранены единые условия задания.")

    heading(doc, "КОЭФФИЦИЕНТЫ И РАЗРЕЖЕННОСТЬ")
    paragraph(doc, "Ниже приведены коэффициенты стандартизованных предикторов при выбранных параметрах. Свободный член не включён в нормы и анализ разреженности. Коэффициенты в исходном масштабе дополнительно сохранены в coefficients.csv.")
    for number, (name, features) in enumerate(MODELS.items(), 1):
        caption(doc, f"Таблица {number+3} — Коэффициенты модели {number}")
        sub = coeff[coeff.model == name]
        rows = [[SHORT[f]] + [f"{sub[(sub.feature == f) & (sub.variant == v)].iloc[0].standardized_weight:.5f}" for v in VARIANTS]
                for f in features]
        table(doc, ["Признак", "Без штрафа", "L1", "L2", "ElasticNet"], rows, [5.5, 2.875, 2.875, 2.875, 2.875], size=10)
    paragraph(doc, "При выбранных параметрах строго нулевых коэффициентов нет. Близким к нулю считается |w| ≤ 0,1 в стандартизованной модели. По этому критерию взаимодействие модели 3 близко к нулю без штрафа (0,04485), с L1 (0,04801) и с ElasticNet (0,07600). С L2 его вес равен 0,18569.")
    paragraph(doc, "L2 снижает общую евклидову норму весов модели 3 с 19,03393 до 18,92267, но отдельный вес взаимодействия увеличивается. Это возможно из-за корреляции синтетического признака с исходными: штраф перераспределяет вклад между ними и не обязан уменьшать каждый вес отдельно.")
    strong = grid[(grid["lambda"] == 10) & (grid.variant == "L1")]
    if len(strong) == 3:
        paragraph(doc, "В исследовании при λ = 10 L1 создаёт 2, 1 и 4 строгих нуля соответственно в моделях 1, 2 и 3. В полной модели обнулены часы учёбы, сон, практика и внеучебная активность. Взаимодействие сохраняется. Такая разреженность получена ценой большой валидационной ошибки, поэтому эти параметры не выбраны для итоговых моделей.")

    heading(doc, "ЗАКЛЮЧЕНИЕ И ВОСПРОИЗВЕДЕНИЕ")
    paragraph(doc, "Реализованы три линейные модели и четыре режима обучения: без регуляризации, L1, L2 и ElasticNet. Используются собственные вычисления на NumPy/Pandas. Готовые регрессоры и оптимизаторы не применяются. Свободный член исключён из штрафа, а обучение и сравнение выполнены при одинаковом разбиении и настройках градиентного спуска.")
    paragraph(doc, "Полный набор признаков обеспечивает R² около 0,9886 и RMSE около 2,0644 на тесте. Регуляризация уменьшает нормы весов, но в этом эксперименте её влияние на прогноз невелико. При недостаточных признаках она не устраняет недообучение. Чрезмерный штраф резко ухудшает качество. Параметры следует выбирать по валидационной ошибке, а не по желанию получить больше нулевых коэффициентов.")
    paragraph(doc, "L1 и ElasticNet способны создавать строгие нули. При выбранных слабых штрафах нулей нет, но синтетический признак полной модели остаётся близким к нулю в трёх вариантах. L2 не создаёт строгих нулей и может перераспределять коэффициенты коррелирующих признаков, уменьшая их общую норму.")
    paragraph(doc, "Обучение реализовано в функции fit файла regularization.py. Графики строит отдельный файл graphics.py. Программа сохраняет сетку параметров, общую таблицу метрик, коэффициенты, прогнозы, разбиение, параметры подготовки, модели и историю всех итераций. Для каждой модели имеются четыре отдельных требуемых графика и объединённый рисунок. Данные читаются из соседнего каталога lab1. Сетевое подключение при запуске эксперимента не требуется.")
    paragraph(doc, "Запуск из каталога Labs/lab3 после установки Python 3.10 или новее:")
    for line in ["python -m pip install -r requirements.txt", "python regularization.py",
                 "python -m unittest -v test_regularization.py"]:
        paragraph(doc, line, "Lab Code")
    paragraph(doc, "Восемь независимых проверок успешно пройдены. Результаты сопоставлены с решением МНК, аналитическим решением L2 и точным решением L1/ElasticNet на ортогональных признаках. Дополнительно проверены условия оптимальности L1/ElasticNet, отсутствие штрафа для свободного члена, корректность историй, заполнение пропусков по train и согласованность сохранённых результатов.")
    paragraph(doc, "Контрольная сумма исходного CSV (SHA-256): " + s["dataset_sha256"] + ". Это позволяет проверить, что повторный запуск использует ту же версию данных.").alignment = WD_ALIGN_PARAGRAPH.LEFT
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)
    QA.mkdir(exist_ok=True)
    authored = QA / "authored.docx"
    doc.save(authored)
    replace_package(authored, FINAL)
    print(FINAL)


def finalize(refreshed):
    """После обновления TOC в Word возвращаем защищённые части шаблона."""
    staged = QA / "field-refreshed-preserved.docx"
    replace_package(refreshed, staged)
    with ZipFile(staged) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    with ZipFile(QA / "authored.docx") as z:
        authored_body = etree.fromstring(z.read("word/document.xml")).find("w:body", NS)
    root = etree.fromstring(parts["word/document.xml"])
    body = root.find("w:body", NS)
    # Титульный лист, включая заполненную тему, сохраняется без переразметки Word.
    for index in range(26):
        body.replace(body[index], deepcopy(authored_body[index]))
    section = body.find("w:sectPr", NS)
    body.replace(section, deepcopy(authored_body.find("w:sectPr", NS)))
    # Word переупорядочивает rId при сохранении. Возвращаем исходные номера
    # отношений, выделяем новые только для добавленных ресурсов и исправляем
    # ссылки на изображения в тексте. Остальные части пакета не затрагиваем.
    with ZipFile(REFERENCE) as original:
        baseline_rels = etree.fromstring(original.read("word/_rels/document.xml.rels"))
    refreshed_rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    identity = lambda rel: (rel.get("Type"), rel.get("Target"), rel.get("TargetMode"))
    baseline_ids = {identity(rel): rel.get("Id") for rel in baseline_rels}
    used_ids = set(baseline_ids.values())
    remap = {}
    next_id = 1
    for rel in refreshed_rels:
        if identity(rel) in baseline_ids:
            new_id = baseline_ids[identity(rel)]
        else:
            while f"rId{next_id}" in used_ids:
                next_id += 1
            new_id = f"rId{next_id}"
            used_ids.add(new_id)
            added = deepcopy(rel)
            added.set("Id", new_id)
            baseline_rels.append(added)
        remap[rel.get("Id")] = new_id
    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for node in root.iter():
        for attr, value in list(node.attrib.items()):
            if attr.startswith("{" + relationship_ns + "}") and value in remap:
                node.set(attr, remap[value])
    parts["word/_rels/document.xml.rels"] = etree.tostring(baseline_rels, xml_declaration=True,
                                                          encoding="UTF-8", standalone=True)
    parts["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(FINAL, "w", ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    with ZipFile(REFERENCE) as original, ZipFile(FINAL) as final:
        for name in original.namelist():
            if name not in EDITABLE:
                assert original.read(name) == final.read(name), name
    print("Проверка неизменяемых частей пройдена:", FINAL)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--finalize", type=Path)
    args = parser.parse_args()
    finalize(args.finalize) if args.finalize else create_report()
