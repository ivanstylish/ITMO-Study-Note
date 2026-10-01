from pathlib import Path
from itertools import product
import importlib.util
import json
import numpy as np
from docx import Document
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('calculation', ROOT / 'lab2_variant15.py')
model = importlib.util.module_from_spec(spec)
spec.loader.exec_module(model)
doc = Document(ROOT / '3310_Чжун_Цзяцзюнь_УИР2.docx')
results = [model.system1(.5), model.system2(.5)]
checks = []

def check(label, condition):
    assert condition, label
    checks.append(label)

def number(text):
    return 0.0 if text == '·' else float(text.replace(',', '.'))

# Reconstruct every generator independently, including event multiplicities.
states1 = list(product((0, 1), repeat=3))
states1 = [(*s, 0) for s in states1] + [(1, 1, 1, q) for q in range(1, 4)]
q1 = np.zeros((11, 11))
for row, s in enumerate(states1):
    for col, t in enumerate(states1):
        if row == col:
            continue
        idle_weights = sum(r for x, r in zip(s[:3], [.7, .2, .1]) if x == 0)
        for i in range(3):
            arrival = list(s)
            arrival[i] = 1
            if not s[i] and tuple(arrival) == t:
                q1[row, col] += .5 * [.7, .2, .1][i] / idle_weights
            service = list(s)
            if s[i]:
                if s[3]:
                    service[3] -= 1
                else:
                    service[i] = 0
                if tuple(service) == t:
                    q1[row, col] += .1
        if s[:3] == (1, 1, 1) and s[3] < 3 and t == (*s[:3], s[3] + 1):
            q1[row, col] += .5
np.fill_diagonal(q1, -q1.sum(axis=1))

def birth_death(size, arrival, service):
    q = np.zeros((size, size))
    for i in range(size - 1):
        q[i, i+1] = arrival
        q[i+1, i] = service
    np.fill_diagonal(q, -q.sum(axis=1))
    return q

alpha = (1 + np.sqrt(39 / 89)) / 2
mu_a, mu_b = alpha / 5, (1-alpha) / 5
phase = np.array([[-.05, .05*alpha, .05*(1-alpha)],
                  [mu_a, -mu_a, 0], [mu_b, 0, -mu_b]])
q2 = (np.kron(birth_death(4, .35, .1), np.eye(9)) +
      np.kron(np.eye(4), np.kron(birth_death(3, .1, .1), np.eye(3))) +
      np.kron(np.eye(12), phase))

for j, (q, result) in enumerate(zip([q1, q2], results), start=1):
    check(f'Q{j}: independent event construction', np.max(abs(q - result['q'])) < 1e-14)
    check(f'Q{j}: rows sum to zero', np.max(abs(q.sum(axis=1))) < 1e-14)
    check(f'Q{j}: off-diagonal rates nonnegative', (q - np.diag(np.diag(q))).min() >= 0)
    a = np.vstack([q.T, np.ones(len(q))])
    pi, *_ = np.linalg.lstsq(a, np.r_[np.zeros(len(q)), 1], rcond=None)
    check(f'pi{j}: independent overdetermined solution', np.max(abs(pi-result['pi'])) < 1e-13)
    csv = np.loadtxt(ROOT / 'Calculation_data' / f'system{j}_Q.csv', delimiter=',')
    check(f'Q{j}: CSV values', np.max(abs(csv-q)) < 1e-11)

printed_q1 = np.array([[number(c.text) for c in row.cells[1:]] for row in doc.tables[2].rows[1:]])
printed_q2 = np.hstack([np.array([[number(c.text) for c in row.cells[1:]]
                               for row in doc.tables[k].rows[1:]]) for k in (4, 5, 6)])
check('Word: all 121 cells of Q1', np.max(abs(printed_q1-q1)) <= .500001e-6)
check('Word: all 1296 cells of Q2', np.max(abs(printed_q2-q2)) <= .500001e-6)
for j, result in enumerate(results):
    values = np.array([number(row.cells[2+2*j].text) for row in doc.tables[7].rows[1:len(result['pi'])+1]])
    check(f'Word: all stationary probabilities of system {j+1}', np.max(abs(values-result['pi'])) <= .500001e-6)

busy1 = results[0]['busy']
x1 = busy1 / 10
l1 = np.full(3, results[0]['L']/3)
w1 = .5 * results[0]['wait_num'] / (3*x1)
for table, expected in [(9, [np.array([3.5,1,.5]), busy1, l1, busy1+l1]),
                        (10, [w1, w1+10, np.array([np.nan]*3), x1])]:
    for group, vals in enumerate(expected):
        for i, val in enumerate(vals):
            if np.isfinite(val):
                printed = number(doc.tables[table].rows[1+group*4+i].cells[3].text.rstrip('*'))
                check(f'Word: system1 table{table} group{group} device{i+1}', abs(printed-val) <= .500001e-6)
    vals2 = ([np.array([3.5,1,.5]), results[1]['busy'], results[1]['queues'],
              results[1]['busy']+results[1]['queues']] if table == 9 else
             [results[1]['wait_i'], results[1]['wait_i']+10, results[1]['loss_i'], results[1]['x_i']])
    for group, vals in enumerate(vals2):
        for i, val in enumerate(vals):
            printed = number(doc.tables[table].rows[1+group*4+i].cells[4].text)
            check(f'Word: system2 table{table} group{group} device{i+1}', abs(printed-val) <= .500001e-6)

for row in doc.tables[12].rows[1:]:
    lam = number(row.cells[0].text)
    a, b = model.system1(lam), model.system2(lam)
    for cell, value in zip(row.cells[1:], [a['L'], b['L'], a['X'], b['X']]):
        check(f'Word: sweep lambda={lam} value={value:.6f}', abs(number(cell.text)-value) <= .500001e-6)
check('H distribution: mean 10s', abs(alpha/mu_a+(1-alpha)/mu_b-10) < 1e-13)
second_moment = 2*alpha/mu_a**2 + 2*(1-alpha)/mu_b**2
check('H distribution: variance 256s^2 and CV 1.6', abs(second_moment-356) < 1e-11)

# Prepared inputs, not a claim that MARK has been run or can import these formats.
inputs = ROOT / 'Calculation_data' / 'MARK_input'
inputs.mkdir(exist_ok=True)
for j, (q, result) in enumerate(zip([q1,q2], results), start=1):
    lines = ['source\ttarget\trate_per_second']
    lines += [f'S{i}\tS{k}\t{q[i,k]:.16g}' for i in range(len(q)) for k in range(len(q)) if i != k and q[i,k] > 0]
    (inputs / f'system{j}_transitions.tsv').write_text('\n'.join(lines)+'\n', encoding='utf-8')
    (inputs / f'system{j}_expected_probabilities.tsv').write_text('state\tprobability\n'+'\n'.join(f'S{i}\t{p:.16g}' for i,p in enumerate(result['pi']))+'\n', encoding='utf-8')
    np.savetxt(inputs / f'system{j}_Q_full_precision.csv', q, delimiter=',', fmt='%.16g')

images = ROOT / '_audit' / 'images'
images.mkdir(exist_ok=True)
with ZipFile(ROOT / '3310_Чжун_Цзяцзюнь_УИР2.docx') as z:
    for name in z.namelist():
        if name.startswith('word/media/'):
            (images / Path(name).name).write_bytes(z.read(name))
summary = {'checks_passed': len(checks), 'checks': checks,
           'metrics': [{k: float(r[k]) for k in ['L','N','loss','X','W','T']} for r in results],
           'stationary_residuals': [float(np.max(abs(r['pi']@r['q']))) for r in results],
           'directed_arcs': [int(np.count_nonzero(q-np.diag(np.diag(q)))) for q in [q1,q2]],
           'MARK_executed': False}
(ROOT / '_audit' / 'audit_results.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in summary.items() if k != 'checks'}, ensure_ascii=False, indent=2))
