from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json, re

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
SOURCE = ROOT / '3310_Чжун_Цзяцзюнь_УИР2.docx'
OUTPUT = ROOT / '3310_Чжун_Цзяцзюнь_УИР2_精简版.docx'

def digest(path):
    return sha256(path.read_bytes()).hexdigest()

protected = [SOURCE, ROOT/'lab2_variant15.py']
protected += [p for p in (ROOT/'Calculation_data').rglob('*') if p.is_file()]
protected += [p for p in (ROOT/'WinMark').rglob('*') if p.is_file()]
before = {str(p): digest(p) for p in protected}

source = Document(SOURCE)
d = Document(SOURCE)
p = list(d.paragraphs)

def replace(index, text):
    paragraph = p[index]
    props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if props is not None:
        run._r.insert(0, props)

revisions = {
29: 'Построить марковские модели двух СМО с ограниченными накопителями, определить стационарные вероятности и характеристики и выбрать систему по заданному критерию эффективности.',
31: 'N1 = 15, N2 = 15, NG = 10. При чётной сумме N1 + N2 = 30 критерий выбирается по N2: «д» — минимальная суммарная средняя длина очередей.',
34: 'Входной поток пуассоновский; времена обслуживания независимы. Дисциплина очередей — FCFS, без прерываний. Ёмкость накопителя не включает обслуживаемые заявки; общая вместимость каждой СМО — шесть заявок.',
35: 'Для общей очереди системы 1 принято дополнительное правило: свободный прибор выбирается с вероятностью rᵢ/Σсвоб. rⱼ. Если все приборы заняты, заявка ожидает; при заполненном накопителе теряется. Ожидающие заявки не закрепляются за приборами; при наличии очереди приборы не простаивают.',
36: 'В системе 2 потоки к приборам равны λ1 = 0,35; λ2 = 0,10; λ3 = 0,05 с⁻¹. При заполнении выбранного прибора и его накопителя заявка теряется независимо от состояния других приборов.',
37: 'Гиперэкспоненциальное обслуживание H1,6 назначено П3, что допускается условием. Сравнение относится к этой конфигурации.',
38: 'Метод расчёта',
39: 'Стационарные вероятности и характеристики обеих СМО рассчитаны в MARK. Характеристики вычислены через вероятности состояний; соотношения Литтла использованы только для проверки.',
42: 'Три экспоненциальных прибора с μ = 1/b = 0,1 с⁻¹ используют общий накопитель на три места. Система соответствует M/M/3 с общей вместимостью 6.',
46: 's = (x1,x2,x3,q), где xᵢ = 1 означает занятый прибор, q — число ожидающих. При q > 0 все приборы заняты. Нумерация S0–S10 едина для графа, матрицы и таблицы вероятностей.',
48: 'Полная модель содержит 11 состояний. Для суммарных характеристик возможна агрегация по n = x1+x2+x3+q: n = 0,…,6; интенсивности рождения λ и смерти min(n,3)μ. Для индивидуальных показателей используется полная модель.',
52: 'Соединения двунаправленные: до S7 интенсивность вправо указана над дугой, влево — 0,1. На цепочке S7–S10 интенсивности равны 0,5 вправо и 0,3 влево. Пересечения без кружка не являются состояниями. Единица интенсивности — с⁻¹.',
55: 'Строка задаёт исходное состояние, столбец — конечное; Qss = −Σt≠s Qst. Отказы не изменяют состояние и не входят в интенсивность выхода.',
56: 'Стационарное распределение удовлетворяет πQ = 0 и Σπs = 1. Для проверки агрегированной цепи: P(n+1) = P(n)λ/[μ min(n+1,3)] с нормировкой. Все вероятности приведены в таблице 6.',
59: 'П1 и П2 имеют экспоненциальное обслуживание и накопители на 2 и 1 место. П3 имеет обслуживание H1,6 без накопителя. Полные локальные вместимости: K1 = 3, K2 = 2, K3 = 1.',
62: 'Параметр v = 1,6 — коэффициент вариации. Использована двухветвевая H2-параметризация с равными вкладами в среднее: при начале обслуживания выбирается A с вероятностью α или B с вероятностью 1−α; после завершения выбранной ветви заявка уходит.',
66: 's = (n1,n2,h): n1 = 0,…,3; n2 = 0,…,2; h ∈ {0,A,B}. n1 и n2 включают обслуживаемую заявку; n3 = 1(h≠0). h = 0 означает свободный П3, A и B — ветви обслуживания. Все 4×3×3 = 36 состояний приведены в таблице 6.',
67: 'Код c(0)=0, c(A)=1, c(B)=2 задаёт номер k = 9n1 + 3n2 + c(h). Например, S0=(0,0,0), S1=(0,0,A), S35=(3,2,B). Сохранены исходные ёмкости и полный набор состояний.',
68: 'Фаза h определяет интенсивность завершения μA или μB и обеспечивает марковское описание. У П3 без очереди загрузка и вероятность отказа зависят от среднего b; это свойство не распространяется на прибор с очередью.',
73: 'Соединения двунаправленные; метки a, b, cA, cB раскрыты в легенде. Пересечения без кружка не являются состояниями. Переходов A↔B нет; отказ не изменяет состояние.',
84: 'Вероятности MARK округлены до шести знаков. Для сохранённых значений |Σπ−1|<10⁻⁶, max|πQ|<10⁻⁷; для высокоточного контрольного решения max|πQ|<10⁻¹².',
87: 'Ui — вероятность занятости; Li — среднее число ожидающих; Ni — среднее число заявок у прибора; Xi — производительность. ΣUi — среднее число занятых приборов; средняя загрузка одного прибора равна ΣUi/3.',
90: 'Для системы 1 P(n)=Σs:Σx+q=n πs. Принятая заявка при n≥3 ожидает n−2 завершений со средним интервалом 1/(3μ). Для пуассоновского входа вероятности состояний при поступлении равны стационарным.',
92: 'Для системы 2 uᵢ(n)=Pr(nᵢ=n), K1=3, K2=2. Принятая заявка, заставшая n заявок у экспоненциального прибора i, ожидает в среднем nb. Формула ниже относится к i=1,2; у П3 W3=0.',
94: 'Времена относятся к принятым заявкам; системные средние взвешиваются по Xi/ΣXi. Равенства Lq=XW и N=XT используются только для проверки.',
95: 'Индивидуальные показатели общей очереди',
96: 'В системе 1 Li=Lq/3, Ni=Ui+Li: при одинаковых экспоненциальных скоростях вероятность будущего прибора для ожидающей заявки равна 1/3 и не зависит от времени ожидания. Поэтому Wi=λD/(3Xi), Ti=Wi+b; отказ относится ко всей СМО.',
97: 'Нагрузки λrᵢb системы 1 являются номинальными весами выбора, а не фиксированными входными потоками; в таблицах они отмечены звёздочкой.',
100: 'Средние загрузки одного прибора ΣUi/3 равны 0,962781 и 0,661076 соответственно.',
101: 'Нагрузка измеряется в эрлангах; Ui — доля времени, L и N — среднее число заявок. Простой прибора ηᵢ=1−Ui; средний простой системы равен 1−ΣUi/3.',
104: 'Времена в строке «Сумм.» взвешены по Xi/X. Прочерк означает отсутствие отдельного отказа по прибору системы 1. Последние знаки чисел MARK зависят от точности программы; таблица 10 содержит высокоточный контрольный расчёт.',
108: 'Зависимости при λ от 0,1 до 1,0 с⁻¹ и неизменных b, rᵢ и ёмкостях приведены на рисунке 5 и в таблице 11. Синяя линия — система 1, коричневая — система 2.',
112: 'С ростом λ длины очередей стремятся к 3, выпуск насыщается, потери возрастают. При малом входе время пребывания стремится к b=10 с.',
113: 'Вывод',
115: 'Построены марковские модели двух СМО и определены стационарные вероятности и показатели их работы.',
116: 'При λ=0,5 с⁻¹ по критерию «д» выбирается система 1: Lq¹=1,925832 против Lq²=1,976939. Общая очередь обеспечивает уменьшение на 0,051107 заявки, или 2,59%, благодаря использованию свободных приборов.',
117: 'Кривые Lq пересекаются между λ=0,5 и 0,6 с⁻¹. При дальнейшем росте входного потока преимущество по длине очереди переходит к системе 2; вывод об оптимальности относится к заданным параметрам.',
}
for index, text in revisions.items():
    replace(index, text)

# Remove implementation and software-operation sections, keeping mathematical checks.
removed = [114] + list(range(118,136))
for index in removed:
    node=p[index]._p
    node.getparent().remove(node)

# Preserve the original model-page groups so diagrams and matrices remain intact.
# The short conclusion can follow the comparative analysis on the same page.
p[113].paragraph_format.page_break_before = False
for index in (113,115,116):
    p[index].paragraph_format.keep_with_next = True
p[117].paragraph_format.keep_with_next = False

# Strict preservation checks: no numerical tables, diagrams, equations, state IDs,
# section page sizes, fonts, cover or source data are rewritten.
assert len(source.tables)==len(d.tables)==13
assert all(a._tbl.xml==b._tbl.xml for a,b in zip(source.tables,d.tables))
assert len(source.inline_shapes)==len(d.inline_shapes)==5
for a,b in zip(source.inline_shapes,d.inline_shapes):
    assert a._inline.xml==b._inline.xml
    rid=a._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
    assert source.part.related_parts[rid].blob==d.part.related_parts[rid].blob
assert [etree.tostring(et) for et in source._element.xpath('.//m:oMath')]==[etree.tostring(et) for et in d._element.xpath('.//m:oMath')]
assert all(a.xml==b.xml for a,b in zip(list(source._element.body)[:27],list(d._element.body)[:27]))
assert [s._sectPr.xml for s in source.sections]==[s._sectPr.xml for s in d.sections]
body='\n'.join(x.text for x in d.paragraphs)
for forbidden in ['Python','NumPy','np.linalg','lab2_variant15.py','CSV','HTML','2003','1.0.0.1','Программа расчёта','Порядок выполнения и точность','Входные данные и сохранённые результаты']:
    assert forbidden not in body, forbidden
for result in ['1,925832','1,976939','0,051107','2,59%']:
    assert result in body
assert len(d.tables[2].rows)==12 and len(d.tables[2].columns)==12
assert all(len(d.tables[i].rows)==37 and len(d.tables[i].columns)==13 for i in [4,5,6])

d.save(OUTPUT)
assert all(digest(Path(path))==old for path,old in before.items())
audit={
 'source':str(SOURCE),'output':str(OUTPUT),
 'source_paragraphs':len(source.paragraphs),'output_paragraphs':len(d.paragraphs),
 'source_body_characters':sum(len(x.text) for x in source.paragraphs[27:]),
 'output_body_characters':sum(len(x.text) for x in d.paragraphs[27:]),
 'tables_unchanged':13,'images_unchanged':5,
 'native_math_objects_unchanged':len(d._element.xpath('.//m:oMath')),
 'protected_files_unchanged':len(before),
 'protected_file_hashes':before,
 'removed_paragraph_indices':removed,
 'issues_to_confirm':['Cover uses 15/10 whereas body specifies 15/15/10; cover retained unchanged.','Common-queue dispatcher is an explicit assumption of the existing model; retained, not reinterpreted.'],
}
(WORK/'preservation_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in audit.items() if k not in ['protected_file_hashes']},ensure_ascii=False,indent=2))
