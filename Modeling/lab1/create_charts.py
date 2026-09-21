"""读取实验 Excel 中的变体 115，拟合 H₂ 并生成六张实验图。

运行：python create_charts.py
依赖：python -m pip install numpy openpyxl matplotlib
可选：python create_charts.py --source 数据.xlsx --output 输出目录
"""

from pathlib import Path
import argparse
import json
import os
import sys

ROOT = Path(__file__).resolve().parent
# 将字体缓存放在实验临时目录，避免受限环境下写入用户配置目录失败。
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "tmp" / "matplotlib_cache"))
# 本机安装的绘图依赖放在实验目录，其他电脑正常 pip 安装后也能运行。
if (ROOT / ".python_deps").is_dir():
    sys.path.insert(0, str(ROOT / ".python_deps"))

import matplotlib
matplotlib.use("Agg")  # 直接保存图片，无需弹出窗口，适合批量生成。
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

SEED = 11501
BLUE, RED, GREEN = "#2E5EAA", "#C95A43", "#1B6B4A"


def load_data(source):
    """按标题 115 定位数据，不是读取第 115 个物理列。"""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        ws = wb["101-150"]
        columns = [cell.column for cell in ws[1] if cell.value == 115]
        if len(columns) != 1:
            raise ValueError("工作表首行必须有唯一的变体标题 115")
        # 一次迭代读取第 2 至 301 行，避免只读模式下反复扫描工作表。
        values = [r[0] for r in ws.iter_rows(
            min_row=2, max_row=301, min_col=columns[0],
            max_col=columns[0], values_only=True)]
    finally:
        wb.close()
    if any(not isinstance(v, (float, int)) or isinstance(v, bool) for v in values):
        raise ValueError("300 个观测值中存在空值或非数值")
    x = np.asarray(values, dtype=float)
    if len(x) != 300 or not np.isfinite(x).all() or (x < 0).any():
        raise ValueError("预期读取 300 个有限、非负观测值")
    return x


def fit_and_generate(x):
    # 均值 x̄=Σx/n；无偏方差 s²=Σ(x-x̄)²/(n-1)，故使用 ddof=1。
    mean = float(x.mean())
    variance = float(x.var(ddof=1))
    cv = float(np.sqrt(variance) / mean)
    if cv <= 1:
        raise ValueError("当前 H₂ 矩匹配要求变异系数 v>1")
    # H₂ 是指数分布的混合；两个矩不足以唯一确定三个参数，额外选 q。
    # q 必须严格小于 2/(1+v²)，以保证第二个分量均值 t2>0。
    q = min(0.3, 0.8 * 2 / (1 + cv**2))
    t1 = mean * (1 + np.sqrt((1-q) * (cv**2-1) / (2*q)))
    t2 = mean * (1 - np.sqrt(q * (cv**2-1) / (2*(1-q))))
    # 理论 E[X]=q*t1+(1-q)*t2；Var(X)=2[q*t1²+(1-q)*t2²]-E[X]²。
    np.testing.assert_allclose(q*t1+(1-q)*t2, mean)
    np.testing.assert_allclose(2*(q*t1**2+(1-q)*t2**2)-mean**2, variance)
    # 显式使用 PCG64，并保持先整批 U1、再整批 U2 的调用顺序，与原报告一致。
    rng = np.random.Generator(np.random.PCG64(SEED))
    u1, u2 = rng.random(len(x)), rng.random(len(x))
    # U1 选分量，U2 做逆变换 X=-t*ln(1-U2)。log1p 提高小数值计算精度。
    generated = np.where(u1 < q, t1, t2) * (-np.log1p(-u2))
    return generated, dict(mean=mean, variance=variance, cv=cv,
                          q=q, t1=float(t1), t2=float(t2))


def acf(x):
    # Pearson 自相关：两个截取序列分别减去各自的均值，与原分析代码一致。
    # 不同于统一减全样本均值的另一种 ACF 定义；检查滞后 1 至 10。
    return np.array([np.corrcoef(x[:-k], x[k:])[0, 1] for k in range(1, 11)])


def axes(title, xlabel, ylabel):
    fig, ax = plt.subplots(figsize=(10, 5), layout="constrained")
    ax.set(title=title, xlabel=xlabel, ylabel=ylabel)
    ax.grid(alpha=0.25)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    return fig, ax


def save(fig, out, name):
    fig.savefig(out / name, dpi=200)
    plt.close(fig)


def plot_sequence(x, out, name, title, color):
    # 不排序，保留观测顺序，才能观察趋势、周期和异常峰值。
    fig, ax = axes(title, "Номер наблюдения", "Значение")
    ax.plot(np.arange(1, len(x)+1), x, color=color, linewidth=1)
    ax.set_xlim(1, len(x))
    ax.set_ylim(bottom=0)
    save(fig, out, name)


def plot_acf(values, n, out, name, title, color):
    fig, ax = axes(title, "Сдвиг", "Коэффициент автокорреляции")
    # 白噪声下单个滞后的近似 95% 界限；不是覆盖所有滞后的同时置信带。
    bound = 1.96 / np.sqrt(n)
    ax.axhspan(-bound, bound, color="#DCECF7", label="95%: ±1.96/√n")
    ax.axhline(0, color="gray", linewidth=0.8)
    ax.plot(range(1, 11), values, "o-", color=color, linewidth=1.5)
    ax.set_xticks(range(1, 11))
    ax.set_ylim(-0.2, 0.2)
    ax.legend(loc="lower right")
    save(fig, out, name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path,
                        default=ROOT / "!_УИР1_Варианты_с — копия.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "output" / "python_charts")
    args = parser.parse_args()
    x = load_data(args.source)
    y, params = fit_and_generate(x)
    rx, ry = acf(x), acf(y)
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    # DejaVu Sans 随 Matplotlib 提供，支持图中俄文；源代码注释为中文。
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11,
                         "axes.titlesize": 14, "savefig.facecolor": "white"})
    plot_sequence(x, out, "01_original_sequence.png", "Исходная последовательность", BLUE)
    plot_acf(rx, len(x), out, "02_original_acf.png", "Автокорреляция исходной последовательности", BLUE)

    # 19 个边界形成 18 个等宽区间；最后一组包含右端点，频数之和为 n。
    edges = np.linspace(x.min(), x.max(), 19)
    fig, ax = axes("Гистограмма исходной последовательности", "Значение", "Частота")
    counts, _, _ = ax.hist(x, bins=edges, color=BLUE, edgecolor="white")
    assert counts.sum() == len(x)
    save(fig, out, "03_original_histogram.png")
    plot_sequence(y, out, "04_generated_sequence.png", "Сгенерированная последовательность", RED)
    plot_acf(ry, len(y), out, "05_generated_acf.png", "Автокорреляция сгенерированной последовательности", RED)

    # 比较时统一两组的分组范围；density=True 实现柱高=频数/(样本量×组距)。
    # 密度直方图总面积为 1，才可与理论概率密度放在同一坐标系比较。
    upper = np.ceil(max(x.max(), y.max()) / 50) * 50
    common_edges = np.linspace(0, upper, 19)
    fig, ax = axes("Сравнение распределений", "Значение", "Плотность")
    ax.hist(x, bins=common_edges, density=True, alpha=0.5, color=BLUE, label="Исходная ЧП")
    ax.hist(y, bins=common_edges, density=True, histtype="step", linewidth=1.8,
            color=RED, label="Сгенерированная ЧП")
    grid = np.linspace(0, upper, 500)
    q, t1, t2 = params["q"], params["t1"], params["t2"]
    # H₂ 密度 f(x)=q/t1*exp(-x/t1)+(1-q)/t2*exp(-x/t2)，x≥0。
    # t1、t2 为分量均值，速率参数分别为 1/t1、1/t2。
    density = q/t1*np.exp(-grid/t1) + (1-q)/t2*np.exp(-grid/t2)
    ax.plot(grid, density, color=GREEN, linewidth=2, label="Плотность H₂")
    ax.set_xlim(0, upper)
    ax.legend()
    save(fig, out, "06_distribution_comparison.png")

    # 一并保存本次实际使用的数据，便于复核图片、参数与旧报告的一致性。
    result = dict(variant=115, seed=SEED, source=str(args.source.resolve()),
                  original=x.tolist(), generated=y.tolist(), parameters=params,
                  original_acf=rx.tolist(), generated_acf=ry.tolist(),
                  histogram_edges=edges.tolist(), histogram_counts=counts.tolist())
    (out / "plot_data.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"已生成六张图：{out.resolve()}")
    print(f"原始均值={x.mean():.5f}，生成均值={y.mean():.5f}，变异系数={params['cv']:.5f}")


if __name__ == "__main__":
    main()
