from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


ROOT = Path(r"D:\Study-Note\Modeling\lab1")
WORK = ROOT / "tmp" / "analysis"
REFERENCE_COPY = WORK / "template_current.docx"
RESULTS = WORK / "variant115" / "results.json"
OUT_DIR = ROOT / "output"
FINAL = OUT_DIR / "УИР1_вариант_115_Чжун_Цзяцзюнь.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

FONT = "Times New Roman"
CHINESE_FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(75, 85, 99)


def set_run_font(run, size=14, bold=None, italic=None, color=BLACK, east_asia=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia or FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.0, keep=False):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_together = keep
    return p


def add_paragraph(doc, text="", size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  before=0, after=0, line=1.0, keep=False, color=BLACK):
    p = doc.add_paragraph()
    format_paragraph(p, align, before, after, line, keep)
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_chinese(doc, text, after=6):
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 2, after, 1.0, True)
    run = p.add_run("中文说明：" + text)
    set_run_font(run, size=10.5, italic=True, color=GRAY, east_asia=CHINESE_FONT)
    return p


def add_heading(doc, text, page_break=False):
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, 8, 6, 1.0, True)
    p.paragraph_format.page_break_before = page_break
    set_run_font(p.add_run(text), size=16, bold=True)
    return p


def add_stage_heading(doc, number, title, page_break=False):
    return add_heading(doc, f"Этап {number} {title}", page_break=page_break)


def add_conclusion(doc, ru, zh):
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 5, 0, 1.0, True)
    set_run_font(p.add_run("Вывод: "), size=14, bold=True)
    set_run_font(p.add_run(ru), size=14)
    add_chinese(doc, zh, after=7)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=45, start=55, bottom=45, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="5F6368", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLACK):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    format_paragraph(p, align, 0, 0, 1.0, True)
    set_run_font(p.add_run(str(text)), size=size, bold=bold, color=color)


def set_col_width(cell, width_inches):
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_table(table, widths, header_rows=1, font_size=9):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if r_idx < header_rows:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            if c_idx < len(widths):
                set_col_width(cell, widths[c_idx])
            if r_idx < header_rows:
                shade_cell(cell, "E7E6E6")
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=r_idx < header_rows)


def fmt_num(v, digits=4):
    return f"{v:.{digits}f}".replace(".", ",")


def fmt_pct(v, digits=2):
    if abs(v) < 0.5 * 10 ** (-digits):
        v = 0.0
    return f"{v:+.{digits}f}%".replace(".", ",")


def add_formula(doc, mathml, width_hint=None):
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    node = etree.fromstring(mathml.encode("utf-8"))
    omml = transform(node).getroot()
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 5, 1.0, True)
    p._p.append(omml)
    return p


def math_doc(inner):
    return f'<math xmlns="http://www.w3.org/1998/Math/MathML" display="block"><mrow>{inner}</mrow></math>'


def mfrac(num, den):
    return f"<mfrac><mrow>{num}</mrow><mrow>{den}</mrow></mfrac>"


def msqrt(x):
    return f"<msqrt><mrow>{x}</mrow></msqrt>"


def sum_i(upper="n"):
    return f"<munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mi>{upper}</mi></munderover>"


def add_picture(doc, filename, caption, width=7.0):
    p = doc.add_paragraph()
    format_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 1, 1.0, True)
    p.add_run().add_picture(str(WORK / "variant115" / filename), width=Inches(width))
    cp = add_paragraph(doc, caption, size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=5, keep=True)
    return p, cp


def add_form_table(doc, stats, relative_key, ci_relative_key, title, header_note=None):
    if header_note:
        add_paragraph(doc, header_note, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=3, keep=True)
    table = doc.add_table(rows=2, cols=8)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(0, 7))
    set_cell(table.cell(0, 0), "Характеристика", 9, True)
    set_cell(table.cell(0, 1), "Строка", 9, True)
    set_cell(table.cell(0, 2), "Количество случайных величин", 9, True)
    for j, n in enumerate(data["sample_sizes"], start=2):
        set_cell(table.cell(1, j), str(n), 9, True)

    row_specs = [
        ("Мат. ожидание", "mean", 5, False),
        ("Дов. инт. 0,90", "0.9", 4, True),
        ("Дов. инт. 0,95", "0.95", 4, True),
        ("Дов. инт. 0,99", "0.99", 4, True),
        ("Дисперсия", "variance", 4, False),
        ("С. к. о.", "std", 4, False),
        ("К-т вариации", "cv", 5, False),
    ]
    for label, key, digits, is_ci in row_specs:
        value_row = table.add_row()
        pct_row = table.add_row()
        value_row.cells[0].merge(pct_row.cells[0])
        set_cell(value_row.cells[0], label, 8.5, True)
        set_cell(value_row.cells[1], "Знач.", 8.5)
        set_cell(pct_row.cells[1], "Δ, %", 8.5)
        for j, n in enumerate(data["sample_sizes"], start=2):
            s = stats[str(n)]
            if is_ci:
                set_cell(value_row.cells[j], "±" + fmt_num(s["ci_half"][key], digits), 8.3)
                set_cell(pct_row.cells[j], fmt_pct(s[ci_relative_key][key]), 8.0)
            else:
                set_cell(value_row.cells[j], fmt_num(s[key], digits), 8.3)
                set_cell(pct_row.cells[j], fmt_pct(s[relative_key][key]), 8.0)
    style_table(table, [1.38, 0.55, 0.88, 0.88, 0.88, 0.88, 0.88, 0.88], header_rows=2, font_size=8.3)
    add_paragraph(doc, title, size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, keep=True)
    return table


def add_histogram_tables(doc):
    edges = data["histogram"]["edges"]
    freqs = data["histogram"]["frequencies"]
    for start in (0, 9):
        stop = min(start + 9, 18)
        table = doc.add_table(rows=4, cols=1 + stop - start)
        labels = ["№", "Левая граница", "Правая граница", "Частота"]
        for i, label in enumerate(labels):
            set_cell(table.cell(i, 0), label, 8.5, True, WD_ALIGN_PARAGRAPH.LEFT)
        for j, idx in enumerate(range(start, stop), start=1):
            set_cell(table.cell(0, j), str(idx + 1), 8.5, True)
            set_cell(table.cell(1, j), fmt_num(edges[idx], 2), 8.0)
            set_cell(table.cell(2, j), fmt_num(edges[idx + 1], 2), 8.0)
            set_cell(table.cell(3, j), str(freqs[idx]), 8.5)
        style_table(table, [1.18] + [0.67] * (stop - start), header_rows=1, font_size=8.2)
        if start == 0:
            add_paragraph(doc, "Продолжение интервалов приведено в следующей части таблицы", size=9.5,
                          italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)


def add_acf_comparison(doc):
    table = doc.add_table(rows=4, cols=11)
    headers = ["Сдвиг"] + [str(i) for i in range(1, 11)]
    for j, value in enumerate(headers):
        set_cell(table.cell(0, j), value, 8.5, True)
    labels = ["Исходная ЧП", "Сгенер. ЧП", "Δ, %"]
    rows = [data["original_acf"], data["generated_acf"], data["acf_relative_difference_pct"]]
    for i, (label, values) in enumerate(zip(labels, rows), start=1):
        set_cell(table.cell(i, 0), label, 8.0, True, WD_ALIGN_PARAGRAPH.LEFT)
        for j, value in enumerate(values, start=1):
            text = fmt_pct(value, 1) if i == 3 else fmt_num(value, 4)
            set_cell(table.cell(i, j), text, 7.8)
    style_table(table, [1.18] + [0.60] * 10, header_rows=1, font_size=7.8)
    add_paragraph(doc, "Форма 3 Коэффициенты автокорреляции", size=10.5, italic=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=4)


def add_cover(doc):
    institution = [
        "Федеральное государственное автономное образовательное",
        "учреждение высшего образования",
        "«Национальный исследовательский университет ИТМО»",
        "Факультет программной инженерии и компьютерной техники",
        "Направление подготовки 09.03.04 «Программная инженерия» –",
        "Системное и прикладное программное обеспечение",
    ]
    for idx, text in enumerate(institution):
        add_paragraph(doc, text, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5,
                      after=58 if idx == len(institution) - 1 else 0)
    for text, size in [
        ("Отчёт", 16),
        ("По лабораторной работе №1", 14),
        ("«Обработка результатов измерений:", 14),
        ("статистический анализ числовой последовательности»", 14),
        ("По моделированию", 14),
        ("Вариант: 115", 14),
    ]:
        add_paragraph(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5,
                      after=24 if text == "Вариант: 115" else 0)

    author_lines = [
        ("Выполнили:", True),
        ("студенты 3 курса", False),
        ("Чжун Цзяцзюнь", False),
        ("Группа: P3310", True),
        ("Принял:", True),
        ("Тропченко Андрей Александрович", False),
        ("Отчёт принят «__»_____2026 г.", False),
        ("Оценка: ___________", False),
    ]
    for text, bold in author_lines:
        add_paragraph(doc, text, size=14, bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT, line=1.15)
    city = add_paragraph(doc, "г. Санкт-Петербург, 2026", size=14, align=WD_ALIGN_PARAGRAPH.CENTER,
                         before=32, after=0)
    city.add_run().add_break(WD_BREAK.PAGE)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(REFERENCE_COPY, FINAL)
    doc = Document(str(FINAL))
    clear_body(doc)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(14)
    normal.font.color.rgb = BLACK

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.different_first_page_header_footer = True

    add_cover(doc)

    add_heading(doc, "Задание")
    add_heading(doc, "Цель работы")
    add_paragraph(
        doc,
        "Изучить методы статистической обработки результатов измерений на примере заданной числовой последовательности: оценить числовые моменты, доверительные интервалы и автокорреляцию, выбрать аппроксимирующий закон по двум начальным моментам, реализовать генератор и сравнить исходную и сгенерированную последовательности.",
        size=14, after=5,
    )
    add_heading(doc, "Содержание отчета")
    requirements = [
        "оценки математического ожидания, дисперсии, среднеквадратического отклонения, коэффициента вариации и доверительных полуинтервалов для n = 10, 20, 50, 100, 200 и 300;",
        "график исходной числовой последовательности и анализ ее характера;",
        "автокорреляционный анализ для сдвигов 1-10;",
        "гистограмма распределения частот;",
        "выбор и расчет параметров аппроксимирующего закона;",
        "описание воспроизводимого алгоритма генерации;",
        "сравнение исходной и сгенерированной последовательностей по числовым характеристикам, распределению, автокорреляции и коэффициенту корреляции.",
    ]
    for i, item in enumerate(requirements, start=1):
        p = add_paragraph(doc, size=12.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        set_run_font(p.add_run(f"{i}. "), size=12.5, bold=True)
        set_run_font(p.add_run(item), size=12.5)
    add_chinese(doc, "本报告依次给出原始序列统计量、图形与自相关分析、分布拟合、可复现生成算法，以及原始序列和生成序列的对比结论。", after=4)

    add_heading(doc, "Ход работы", page_break=True)
    add_heading(doc, "Исходные данные")
    add_paragraph(doc, "Источник данных: файл «!_УИР1_Варианты_с — копия.xlsx», лист «101-150», столбец P, ячейки P2:P301. Столбец содержит 300 наблюдений варианта 115.", size=13, after=4)
    add_paragraph(doc, f"Минимальное значение равно {fmt_num(min(data['original']), 5)}, максимальное - {fmt_num(max(data['original']), 5)}.", size=13, after=4)
    add_chinese(doc, "数据取自工作簿“101-150”工作表的 P2:P301，共 300 个观测值，对应第 115 变体。", after=5)

    first10 = data["original"][:10]
    table = doc.add_table(rows=2, cols=6)
    set_cell(table.cell(0, 0), "i", 9, True)
    set_cell(table.cell(1, 0), "xᵢ", 9, True)
    for j in range(5):
        set_cell(table.cell(0, j + 1), str(j + 1), 9, True)
        set_cell(table.cell(1, j + 1), fmt_num(first10[j], 5), 8.5)
    style_table(table, [0.45] + [1.22] * 5, header_rows=1, font_size=8.5)
    table2 = doc.add_table(rows=2, cols=6)
    set_cell(table2.cell(0, 0), "i", 9, True)
    set_cell(table2.cell(1, 0), "xᵢ", 9, True)
    for j in range(5, 10):
        set_cell(table2.cell(0, j - 4), str(j + 1), 9, True)
        set_cell(table2.cell(1, j - 4), fmt_num(first10[j], 5), 8.5)
    style_table(table2, [0.45] + [1.22] * 5, header_rows=1, font_size=8.5)
    add_paragraph(doc, "Первые 10 значений исходной ЧП", size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)

    add_heading(doc, "Методика и пример расчета")
    mean_eq = math_doc(
        '<mover><mi>x</mi><mo>¯</mo></mover><msub><mi></mi><mi>n</mi></msub><mo>=</mo>' +
        mfrac(sum_i() + '<msub><mi>x</mi><mi>i</mi></msub>', '<mi>n</mi>') +
        '<mo>,</mo><mspace width="1em"/><msubsup><mi>s</mi><mi>n</mi><mn>2</mn></msubsup><mo>=</mo>' +
        mfrac(sum_i() + '<msup><mrow><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>', '<mi>n</mi><mo>-</mo><mn>1</mn>')
    )
    add_formula(doc, mean_eq)
    add_formula(doc, math_doc(
        '<msub><mi>s</mi><mi>n</mi></msub><mo>=</mo>' + msqrt('<msubsup><mi>s</mi><mi>n</mi><mn>2</mn></msubsup>') +
        '<mo>,</mo><mspace width="1em"/><msub><mi>v</mi><mi>n</mi></msub><mo>=</mo>' +
        mfrac('<msub><mi>s</mi><mi>n</mi></msub>', '<mover><mi>x</mi><mo>¯</mo></mover><msub><mi></mi><mi>n</mi></msub>')
    ))
    add_formula(doc, math_doc(
        '<msub><mi>ε</mi><mi>p</mi></msub><mo>=</mo><msub><mi>t</mi><mi>p</mi></msub>' +
        mfrac('<msub><mi>s</mi><mi>n</mi></msub>', msqrt('<mi>n</mi>')) +
        '<mo>,</mo><mspace width="1em"/><msub><mi>I</mi><mi>p</mi></msub><mo>=</mo><mo>[</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>-</mo><msub><mi>ε</mi><mi>p</mi></msub><mo>;</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>+</mo><msub><mi>ε</mi><mi>p</mi></msub><mo>]</mo>'
    ))
    add_formula(doc, math_doc(
        '<msub><mi>ΔA</mi><mi>n</mi></msub><mo>=</mo>' +
        mfrac('<msub><mi>A</mi><mi>n</mi></msub><mo>-</mo><msub><mi>A</mi><mn>300</mn></msub>', '<msub><mi>A</mi><mn>300</mn></msub>') +
        '<mo>·</mo><mn>100</mn><mo>%</mo>'
    ))
    n10 = data["original_stats"]["10"]
    sum10 = sum(first10)
    ss10 = sum((x - n10["mean"]) ** 2 for x in first10)
    mean300 = data["original_stats"]["300"]["mean"]
    ci95lo = n10["mean"] - n10["ci_half"]["0.95"]
    ci95hi = n10["mean"] + n10["ci_half"]["0.95"]
    add_paragraph(doc, "Числовая подстановка для первых десяти наблюдений:", size=13, bold=True, after=2)
    calculations = [
        f"Σxᵢ = {fmt_num(sum10, 5)}; x̄₁₀ = {fmt_num(sum10, 5)} / 10 = {fmt_num(n10['mean'], 5)}.",
        f"Σ(xᵢ - x̄₁₀)² = {fmt_num(ss10, 5)}; s₁₀² = {fmt_num(ss10, 5)} / 9 = {fmt_num(n10['variance'], 5)}.",
        f"s₁₀ = √{fmt_num(n10['variance'], 5)} = {fmt_num(n10['std'], 5)}; v₁₀ = {fmt_num(n10['std'], 5)} / {fmt_num(n10['mean'], 5)} = {fmt_num(n10['cv'], 5)}.",
        f"Для p = 0,95: ε = 1,960 · {fmt_num(n10['std'], 5)} / √10 = {fmt_num(n10['ci_half']['0.95'], 5)}; I₀,₉₅ = [{fmt_num(ci95lo, 5)}; {fmt_num(ci95hi, 5)}].",
        f"Отклонение среднего от эталона n = 300: ({fmt_num(n10['mean'], 5)} - {fmt_num(mean300, 5)}) / {fmt_num(mean300, 5)} · 100% = {fmt_pct(n10['rel_vs_300_pct']['mean'], 2)}.",
    ]
    for text in calculations:
        add_paragraph(doc, text, size=12.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    add_chinese(doc, "均值使用算术平均，方差采用 n-1 分母的无偏估计；置信区间表中填写的是半区间 ±ε。相对偏差以 n=300 的结果为基准，因此 n=300 行的偏差必须为 0%。", after=7)

    add_stage_heading(doc, 1, "Статистические характеристики исходной последовательности", page_break=True)
    add_form_table(doc, data["original_stats"], "rel_vs_300_pct", "ci_rel_vs_300_pct",
                   "Форма 1 Характеристики заданной ЧП вариант 115")
    add_conclusion(
        doc,
        f"При увеличении объема выборки оценки стабилизируются около значений x̄ = {fmt_num(mean300, 5)}, s² = {fmt_num(data['original_stats']['300']['variance'], 5)} и v = {fmt_num(data['original_stats']['300']['cv'], 5)}. Доверительный полуинтервал 95% уменьшается до ±{fmt_num(data['original_stats']['300']['ci_half']['0.95'], 5)}, что отражает рост точности оценки среднего.",
        "随着样本量增大，各统计量逐步稳定。n=300 时均值约为 19.59075，95% 置信半区间缩小到 ±3.89518，说明均值估计的精度提高。",
    )

    add_stage_heading(doc, 2, "График исходной числовой последовательности", page_break=True)
    add_picture(doc, "01_original_sequence.png", "Рисунок 1 Исходная числовая последовательность")
    add_conclusion(
        doc,
        "На графике отсутствует устойчивое возрастание, убывание или повторяющийся период. Большинство наблюдений невелики, но встречаются редкие выбросы до 283,17, поэтому последовательность имеет выраженную правостороннюю асимметрию.",
        "序列没有持续上升、下降或明显周期。多数值较小，但存在最高约 283.17 的少量尖峰，因此分布呈明显右偏。",
    )

    add_stage_heading(doc, 3, "Автокорреляционный анализ")
    acf_math = math_doc(
        '<msub><mi>r</mi><mi>k</mi></msub><mo>=</mo>' + mfrac(
            '<munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mrow><mi>n</mi><mo>-</mo><mi>k</mi></mrow></munderover>'
            '<mrow><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow>'
            '<mrow><mo>(</mo><msub><mi>x</mi><mrow><mi>i</mi><mo>+</mo><mi>k</mi></mrow></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow>',
            msqrt(
                '<munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mrow><mi>n</mi><mo>-</mo><mi>k</mi></mrow></munderover><msup><mrow><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>'
                '<munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mrow><mi>n</mi><mo>-</mo><mi>k</mi></mrow></munderover><msup><mrow><mo>(</mo><msub><mi>x</mi><mrow><mi>i</mi><mo>+</mo><mi>k</mi></mrow></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>'
            )
        )
    )
    add_formula(doc, acf_math)
    table = doc.add_table(rows=2, cols=11)
    for j, value in enumerate(["Сдвиг"] + [str(i) for i in range(1, 11)]):
        set_cell(table.cell(0, j), value, 8.5, True)
    set_cell(table.cell(1, 0), "К-т АК", 8.5, True)
    for j, value in enumerate(data["original_acf"], start=1):
        set_cell(table.cell(1, j), fmt_num(value, 4), 8.0)
    style_table(table, [1.05] + [0.61] * 10, header_rows=1, font_size=8.0)
    add_paragraph(doc, "Коэффициенты автокорреляции исходной ЧП", size=10.5, italic=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_picture(doc, "02_original_acf.png", "Рисунок 2 Автокорреляция исходной ЧП")
    add_conclusion(
        doc,
        f"Приближенная 95%-я граница для случайной последовательности составляет ±1,96/√300 = ±{fmt_num(data['acf_95_bound'], 4)}. Девять из десяти коэффициентов находятся внутри границ; r₇ = {fmt_num(data['original_acf'][6], 4)} выходит за верхнюю границу. Устойчивого периодического рисунка нет, однако строгая гипотеза полной независимости по этому набору лагов подтверждена не полностью.",
        "随机序列的近似 95% 界限为 ±0.1132。10 个滞后中有 9 个位于界限内，但第 7 阶自相关为 0.1585，超过上界。因此没有稳定周期，但不能仅凭该结果完全确认所有观测相互独立。",
    )

    add_stage_heading(doc, 4, "Гистограмма распределения частот", page_break=True)
    add_histogram_tables(doc)
    add_picture(doc, "03_original_histogram.png", "Рисунок 3 Гистограмма исходной ЧП")
    first_freq = data["histogram"]["frequencies"][0]
    add_conclusion(
        doc,
        f"В первый интервал от {fmt_num(data['histogram']['edges'][0], 2)} до {fmt_num(data['histogram']['edges'][1], 2)} попало {first_freq} из 300 наблюдений ({fmt_num(first_freq / 3, 1)}%). Длинный правый хвост и коэффициент вариации v = {fmt_num(data['original_stats']['300']['cv'], 5)} > 1 исключают равномерную и обычную экспоненциальную модели и указывают на гиперэкспоненциальный закон.",
        "首个区间包含 236 个观测值，占 78.7%；同时右侧存在长尾，且变异系数 1.75704 大于 1，因此选择超指数分布更合适。",
    )

    add_stage_heading(doc, 5, "Параметры аппроксимирующего закона", page_break=True)
    add_paragraph(doc, f"Для n = 300 получены t = M(X) = {fmt_num(mean300, 6)} и v = {fmt_num(data['original_stats']['300']['cv'], 6)}. Так как v > 1, используется двухфазный гиперэкспоненциальный закон H₂.", size=13, after=4)
    hyper_math = math_doc(
        '<mi>q</mi><mo>≤</mo>' + mfrac('<mn>2</mn>', '<mn>1</mn><mo>+</mo><msup><mi>v</mi><mn>2</mn></msup>') +
        '<mo>,</mo><mspace width="1em"/><msub><mi>t</mi><mn>1</mn></msub><mo>=</mo><mrow><mo>[</mo><mn>1</mn><mo>+</mo>' +
        msqrt(mfrac('<mn>1</mn><mo>-</mo><mi>q</mi>', '<mn>2</mn><mi>q</mi>') + '<mrow><mo>(</mo><msup><mi>v</mi><mn>2</mn></msup><mo>-</mo><mn>1</mn><mo>)</mo></mrow>') +
        '<mo>]</mo></mrow><mi>t</mi><mo>,</mo><mspace width="1em"/><msub><mi>t</mi><mn>2</mn></msub><mo>=</mo><mrow><mo>[</mo><mn>1</mn><mo>-</mo>' +
        msqrt(mfrac('<mi>q</mi>', '<mn>2</mn><mrow><mo>(</mo><mn>1</mn><mo>-</mo><mi>q</mi><mo>)</mo></mrow>') + '<mrow><mo>(</mo><msup><mi>v</mi><mn>2</mn></msup><mo>-</mo><mn>1</mn><mo>)</mo></mrow>') +
        '<mo>]</mo></mrow><mi>t</mi>'
    )
    add_formula(doc, hyper_math)
    params = data["hyperexponential"]
    theoretical_var = 2 * (params["q"] * params["t1"] ** 2 + (1 - params["q"]) * params["t2"] ** 2) - mean300 ** 2
    add_paragraph(doc, f"Допустимая верхняя граница q: 2/(1 + v²) = {fmt_num(params['q_max'], 6)}. Выбираем q = {fmt_num(params['q'], 1)}.", size=13, after=2)
    add_paragraph(doc, f"t₁ = [1 + √((1 - 0,3)/(2·0,3)·(v² - 1))]·t = {fmt_num(params['t1'], 6)}.", size=13, after=2)
    add_paragraph(doc, f"t₂ = [1 - √(0,3/(2·(1 - 0,3))·(v² - 1))]·t = {fmt_num(params['t2'], 6)}.", size=13, after=2)
    add_paragraph(doc, f"Проверка моментов: q·t₁ + (1-q)·t₂ = {fmt_num(params['q']*params['t1']+(1-params['q'])*params['t2'], 6)}; 2[q·t₁² + (1-q)·t₂²] - t² = {fmt_num(theoretical_var, 6)}.", size=12.5, after=3)
    add_chinese(doc, "由于 v>1，采用二相超指数分布 H₂。取 q=0.3，计算得到 t₁=50.16134、t₂=6.48907。代回后理论均值和方差与 n=300 的样本矩一致。", after=5)
    add_conclusion(doc, "По двум начальным моментам исходную ЧП аппроксимируем смесью двух экспоненциальных распределений с вероятностями 0,3 и 0,7 и средними t₁ и t₂.", "拟合模型是两个指数分布的混合：以 0.3 的概率使用均值 t₁，以 0.7 的概率使用均值 t₂。")

    add_stage_heading(doc, 6, "Алгоритм генерации", page_break=True)
    gen_math = math_doc(
        '<mi>X</mi><mo>=</mo><mfenced open="{" close=""><mtable>'
        '<mtr><mtd><mo>-</mo><msub><mi>t</mi><mn>1</mn></msub><mi>ln</mi><mrow><mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>U</mi><mn>2</mn></msub><mo>)</mo></mrow></mtd><mtd><mtext>если </mtext><msub><mi>U</mi><mn>1</mn></msub><mo>&lt;</mo><mi>q</mi></mtd></mtr>'
        '<mtr><mtd><mo>-</mo><msub><mi>t</mi><mn>2</mn></msub><mi>ln</mi><mrow><mo>(</mo><mn>1</mn><mo>-</mo><msub><mi>U</mi><mn>2</mn></msub><mo>)</mo></mrow></mtd><mtd><mtext>если </mtext><msub><mi>U</mi><mn>1</mn></msub><mo>≥</mo><mi>q</mi></mtd></mtr>'
        '</mtable></mfenced><mo>,</mo><mspace width="1em"/><msub><mi>U</mi><mn>1</mn></msub><mo>,</mo><msub><mi>U</mi><mn>2</mn></msub><mo>∼</mo><mi>U</mi><mrow><mo>(</mo><mn>0</mn><mo>,</mo><mn>1</mn><mo>)</mo></mrow>'
    )
    add_formula(doc, gen_math)
    algorithm = [
        "Инициализировать генератор псевдослучайных чисел фиксированным seed = 11501, составленным из номера варианта 115 и номера работы 01.",
        "Для каждого элемента получить независимые U₁ и U₂, равномерно распределенные на (0; 1).",
        "Если U₁ < 0,3, выбрать масштаб t₁ = 50,161340; иначе выбрать t₂ = 6,489068.",
        "Преобразовать U₂ методом обратной функции: X = -t·ln(1-U₂). Повторить 300 раз.",
        "Для защиты в Excel ту же операцию выполняет формула =ЕСЛИ(A2<$F$2;$H$1;$H$2)*(-ЛН(1-B2)), где A2 и B2 содержат U₁ и U₂.",
    ]
    for i, text in enumerate(algorithm, start=1):
        p = add_paragraph(doc, size=12.5, after=2)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.20)
        set_run_font(p.add_run(f"{i}. "), size=12.5, bold=True)
        set_run_font(p.add_run(text), size=12.5)
    add_chinese(doc, "固定随机种子 11501 使本报告中的 300 个生成值能够复现。每个值先按 U₁ 选择一个指数分量，再用 -t·ln(1-U₂) 生成。", after=5)
    add_picture(doc, "04_generated_sequence.png", "Рисунок 4 Сгенерированная числовая последовательность")
    add_conclusion(doc, "Алгоритм формирует независимые значения заданной смеси и воспроизводит результат при повторном запуске с тем же seed.", "该算法按拟合后的混合分布生成独立值；使用相同随机种子时，结果可以完全复现。")

    add_stage_heading(doc, 7, "Сравнение исходной и сгенерированной последовательностей", page_break=True)
    add_form_table(doc, data["generated_stats"], "rel_vs_original_pct", "ci_rel_vs_original_pct",
                   "Форма 2 Характеристики сгенерированной ЧП",
                   "Закон распределения: двухфазный гиперэкспоненциальный H₂, q = 0,3; seed = 11501")
    g300 = data["generated_stats"]["300"]
    orig300 = data["original_stats"]["300"]
    orig_ci95 = (orig300["mean"] - orig300["ci_half"]["0.95"], orig300["mean"] + orig300["ci_half"]["0.95"])
    add_paragraph(doc, f"Для n = 300 отклонения сгенерированной ЧП от исходной составили: среднее {fmt_pct(g300['rel_vs_original_pct']['mean'], 2)}, дисперсия {fmt_pct(g300['rel_vs_original_pct']['variance'], 2)}, СКО {fmt_pct(g300['rel_vs_original_pct']['std'], 2)}, коэффициент вариации {fmt_pct(g300['rel_vs_original_pct']['cv'], 2)}.", size=12.5, after=3)
    add_paragraph(doc, f"Среднее сгенерированной ЧП {fmt_num(g300['mean'], 5)} находится внутри 95%-го доверительного интервала исходного среднего [{fmt_num(orig_ci95[0], 5)}; {fmt_num(orig_ci95[1], 5)}].", size=12.5, after=4)
    add_chinese(doc, "n=300 时，生成序列的均值偏差为 +1.12%，方差偏差为 +10.61%，标准差偏差为 +5.17%，变异系数偏差为 +4.01%。生成均值位于原始均值的 95% 置信区间内。", after=5)
    add_picture(doc, "06_distribution_comparison.png", "Рисунок 5 Сравнение гистограмм и плотности H₂")
    add_conclusion(doc, "Сгенерированная последовательность воспроизводит основной правосторонне-асимметричный характер и длинный хвост исходных данных. Расхождения отдельных столбцов гистограммы ожидаемы для выборки из 300 случайных значений.", "生成序列再现了原始数据的主要右偏特征和长尾。由于样本量只有 300，个别直方图区间存在差异是正常的。")

    add_heading(doc, "Автокорреляция и взаимная корреляция")
    add_acf_comparison(doc)
    add_picture(doc, "05_generated_acf.png", "Рисунок 6 Автокорреляция сгенерированной ЧП")
    add_formula(doc, math_doc(
        '<msub><mi>r</mi><mrow><mi>x</mi><mi>y</mi></mrow></msub><mo>=</mo>' +
        mfrac(
            sum_i() + '<mrow><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mrow><mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>-</mo><mover><mi>y</mi><mo>¯</mo></mover><mo>)</mo></mrow>',
            msqrt(sum_i() + '<msup><mrow><mo>(</mo><msub><mi>x</mi><mi>i</mi></msub><mo>-</mo><mover><mi>x</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>' + sum_i() + '<msup><mrow><mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>-</mo><mover><mi>y</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup>')
        )
    ))
    add_paragraph(doc, f"Коэффициент корреляции между исходной и сгенерированной последовательностями r = {fmt_num(data['pair_correlation'], 6)}.", size=13, bold=True, after=3)
    add_conclusion(
        doc,
        f"Для сгенерированной ЧП девять из десяти коэффициентов автокорреляции лежат внутри границ ±{fmt_num(data['acf_95_bound'], 4)}; r₁₀ = {fmt_num(data['generated_acf'][9], 4)} немного превышает границу. Последовательной зависимости по лагам не наблюдается. Значение r = {fmt_num(data['pair_correlation'], 4)} близко к нулю и подтверждает отсутствие линейной связи между двумя независимо сформированными реализациями. Процентные различия коэффициентов в форме 3 следует трактовать осторожно, поскольку деление на исходные коэффициенты, близкие к нулю, дает большие проценты.",
        "生成序列的 10 个自相关系数中有 9 个位于 ±0.1132 内，第 10 阶略超出；未见连续的滞后依赖。两序列相关系数为 -0.0223，接近 0，符合独立生成的预期。由于原始自相关接近 0，百分比差值可能很大，不能单独作为拟合优劣指标。",
    )

    add_heading(doc, "Выводы", page_break=True)
    final_ru = (
        f"Для варианта 115 обработаны 300 наблюдений. Эталонные оценки составили: математическое ожидание {fmt_num(orig300['mean'], 5)}, "
        f"дисперсия {fmt_num(orig300['variance'], 5)}, среднеквадратическое отклонение {fmt_num(orig300['std'], 5)} и коэффициент вариации {fmt_num(orig300['cv'], 5)}. "
        "График и гистограмма показывают отсутствие тренда и выраженную правостороннюю асимметрию. Автокорреляция в основном мала, хотя коэффициент при сдвиге 7 выходит за приближенную 95%-ю границу, поэтому вывод о полной независимости следует считать осторожным. "
        f"Так как v > 1, по двум начальным моментам выбран двухфазный гиперэкспоненциальный закон с q = 0,3, t₁ = {fmt_num(params['t1'], 5)} и t₂ = {fmt_num(params['t2'], 5)}. "
        f"В воспроизводимой реализации из 300 значений среднее отличается от исходного на {fmt_pct(g300['rel_vs_original_pct']['mean'], 2)}, а коэффициент вариации - на {fmt_pct(g300['rel_vs_original_pct']['cv'], 2)}. "
        f"Среднее сгенерированной ЧП входит в 95%-й доверительный интервал исходного среднего; взаимная корреляция r = {fmt_num(data['pair_correlation'], 4)} близка к нулю. "
        "Следовательно, модель H₂ удовлетворительно воспроизводит первые два момента и общий вид распределения варианта 115, но единичные отклонения автокорреляции не позволяют утверждать абсолютную независимость по всем проверенным лагам."
    )
    add_paragraph(doc, final_ru, size=14, after=5)
    add_chinese(
        doc,
        "第 115 变体的 300 个观测值均值为 19.59075，方差为 1184.84974，标准差为 34.42165，变异系数为 1.75704。由于变异系数大于 1，选用二相超指数分布。生成序列的均值与原始序列仅相差 +1.12%，变异系数相差 +4.01%，且生成均值落在原始均值的 95% 置信区间内。该模型能较好复现前两个矩和整体右偏长尾形态，但原始序列第 7 阶及生成序列第 10 阶自相关略超 95% 界限，因此关于完全独立性的结论应保持谨慎。",
        after=0,
    )

    doc.save(str(FINAL))
    print(FINAL)


if __name__ == "__main__":
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    build()
