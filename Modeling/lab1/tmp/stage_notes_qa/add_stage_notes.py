"""在八个阶段性结论后添加俄语分析，另存报告；原文件不变。"""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "УИР1_вариант_115_Чжун_Цзяцзюнь.docx"
OUTPUT = ROOT / "УИР1_вариант_115_Чжун_Цзяцзюнь_с_пояснениями.docx"

NOTES = [
    (
        "Для таблицы я брал первые 10, 20, 50, 100, 200 и 300 значений. "
        "Среднее считал как сумму, делённую на n, дисперсию — по квадратам отклонений "
        "с делением на n − 1; затем находил СКО и коэффициент вариации. "
        "Приближённый полуинтервал 95% получил, умножив стандартную ошибку среднего на 1,96. "
        "Оценки меняются неравномерно, поэтому я сравнивал их с результатом для всех 300 наблюдений."
    ),
    (
        "Я нанёс на график все 300 значений в том порядке, в котором они записаны в таблице. "
        "Большая часть точек находится внизу, а отдельные крупные значения дают высокие пики; "
        "сами по себе такие пики ещё не означают ошибку измерения. "
        "По этому графику я не вижу устойчивого тренда, а правую асимметрию проверяю дальше по гистограмме."
    ),
    (
        "Здесь я использовал все 300 наблюдений, а k менял от 1 до 10. При k = 1 сравнивал "
        "соседние значения, при k = 2 — значения через одно. В программе для двух отрезков "
        "вычисляется корреляция Пирсона с их собственными средними. Граница относится к "
        "отдельному сдвигу, поэтому один выход за неё не доказывает наличие периода. "
        "При этом считать независимость всех наблюдений установленной я тоже не могу."
    ),
    (
        "Границы групп я получил, разделив диапазон от минимального до максимального значения "
        "на 18 равных частей. Затем подсчитал, сколько наблюдений попало в каждую группу: "
        "для первой получилось 236, или 78,7% выборки. По этим частотам хорошо видно, почему "
        "гистограмма вытянута вправо: малые значения встречаются часто, а большие занимают "
        "широкий диапазон, но встречаются редко."
    ),
    (
        "Я подставил среднее и коэффициент вариации в формулы для двух компонент. "
        "Вероятность q = 0,3 выбрал из допустимого интервала: по двум моментам все три "
        "параметра однозначно не определяются. Значения t₁ и t₂ — это средние компонент, "
        "а не их вероятности. Проверка теоретического среднего и дисперсии показала совпадение "
        "с исходными оценками, но не доказывает совпадение всего распределения."
    ),
    (
        "В программе я сначала получаю два массива равномерных случайных чисел. U₁ выбирает "
        "компоненту, а U₂ подставляется в формулу с логарифмом; их сумма не обязана быть равна "
        "единице. Так получаются 300 новых значений. Фиксированный seed при том же порядке "
        "вызовов генератора позволяет повторить результат. Совпадения отдельных пиков новой "
        "и исходной последовательностей я не ожидаю."
    ),
    (
        "Для сравнения я пересчитал характеристики новой последовательности. Разность новой "
        "и исходной оценок при одинаковом числе наблюдений делил на исходную и умножал на 100%. Например, изменение среднего "
        "с 19,59075 до 19,80944 даёт около 1,12%. Для гистограмм использовал общие интервалы "
        "и нормировку площади до единицы. Зелёную кривую получил сложением двух экспоненциальных "
        "плотностей с весами 0,3 и 0,7; столбцы показывают плотность в среднем по интервалам, "
        "поэтому не обязаны совпадать с кривой в каждой точке."
    ),
    (
        "Здесь я отдельно проверил зависимость внутри сгенерированной последовательности "
        "и связь между двумя последовательностями. Для второго расчёта сопоставил значения "
        "с одинаковыми номерами и получил r = −0,0223. Это говорит о слабой линейной связи, "
        "но само по себе не доказывает независимость. Большие проценты в форме 3 я оцениваю "
        "осторожно: они могут возникать из-за деления на модуль исходного коэффициента, близкого к нулю."
    ),
]

source_bytes = SOURCE.read_bytes()
doc = Document(BytesIO(source_bytes))
original_paragraphs = [p.text for p in doc.paragraphs]
conclusions = [p for p in doc.paragraphs if p.text.strip().startswith("Вывод:")]
assert len(conclusions) == len(NOTES) == 8

for conclusion, text in zip(conclusions, NOTES):
    element = OxmlElement("w:p")
    if conclusion._p.pPr is not None:
        element.append(deepcopy(conclusion._p.pPr))
    conclusion._p.addnext(element)
    paragraph = Paragraph(element, conclusion._parent)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.page_break_before = False
    run = paragraph.add_run(text)
    # 沿用结论正文的字体，避免把“Вывод”标签的粗体带入新增分析。
    body_run = next((r for r in conclusion.runs if r.text.strip() and not r.bold), None)
    if body_run is not None and body_run._r.rPr is not None:
        run._r.insert(0, deepcopy(body_run._r.rPr))
    run.bold = False
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "ru-RU")
    run._r.get_or_add_rPr().append(language)
    conclusion.paragraph_format.keep_with_next = True

# 只改正文 XML，以保留 Word 数学公式、批注及其他包部件。
modified_xml = etree.tostring(doc._element, encoding="UTF-8", xml_declaration=True, standalone=True)
with ZipFile(BytesIO(source_bytes)) as before, ZipFile(OUTPUT, "w") as after:
    for entry in before.infolist():
        data = modified_xml if entry.filename == "word/document.xml" else before.read(entry.filename)
        after.writestr(entry, data)

check = Document(OUTPUT)
assert [p.text for p in check.paragraphs if p.text not in NOTES] == original_paragraphs
assert len(check.tables) == len(doc.tables) == 7
assert len(check.inline_shapes) == len(doc.inline_shapes) == 6
assert all(sum(p.text == note for p in check.paragraphs) == 1 for note in NOTES)
with ZipFile(BytesIO(source_bytes)) as before, ZipFile(OUTPUT) as after:
    assert before.namelist() == after.namelist()
    assert all(before.read(n) == after.read(n) for n in before.namelist() if n != "word/document.xml")
assert SOURCE.read_bytes() == source_bytes
print(f"Added {len(NOTES)} analysis paragraphs: {OUTPUT}")

